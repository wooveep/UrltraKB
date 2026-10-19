"""Non-model work is observable at actual parsing and storage boundaries."""

import json
from zipfile import ZipFile

import pytest
import test_storage_task_entrypoints as task_cases
from stream_transport import StreamTransport
from test_storage_questions import question_knowledge

from ultrakb.application.documents import import_document
from ultrakb.application.execution import ExecutionContext
from ultrakb.application.knowledge_tasks import read_knowledge_task
from ultrakb.application.views import storage_scope

pytestmark = pytest.mark.integration
imported_task = task_cases.imported_task


def test_failed_worksheet_has_a_named_durable_stage_without_losing_other_sheets(tmp_path):
    from openpyxl import Workbook

    root, identity, _ = question_knowledge(tmp_path)
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "VALUE_17"
    book.create_sheet("Beta")["A1"] = "Invalid worksheet XML"
    source = tmp_path / "book.xlsx"
    book.save(source)
    with ZipFile(source) as archive:
        entries = [(entry, archive.read(entry)) for entry in archive.infolist()]
    with ZipFile(source, "w") as archive:
        for entry, payload in entries:
            archive.writestr(
                entry, b"<worksheet broken=" if entry.filename.endswith("sheet2.xml") else payload
            )
    original_size = source.stat().st_size
    context = ExecutionContext(storage_identity=identity)
    responses = [
        {"content": json.dumps({"description": "Value", "content": "VALUE_17"})},
        {"content": json.dumps({"concepts": [], "entities": []})},
    ]
    with StreamTransport(responses) as service:
        (root / ".env").write_text(f"LLM_API_KEY=test-key\nOPENAI_API_BASE={service.url}\n")
        result = import_document(root, source, context=context, scope=storage_scope(root, identity))
    assert result.status == "partial", result.message
    units = {unit.name: unit for unit in result.units}
    assert units["Alpha"].status == "completed" and units["Beta"].status == "failed"
    source.unlink()
    detail = read_knowledge_task(root, context.storage_execution.task_id, limit=128)
    failed = [
        stage
        for stage in detail.stages
        if stage.unit_revision_id == units["Beta"].target_revision_id and stage.state == "failed"
    ]
    assert failed, "The failed worksheet must remain locatable in durable task details"
    assert failed[0].input.unit_label == "Beta"
    assert failed[0].input.name == "book.xlsx"
    assert failed[0].input.original_bytes == original_size
    assert failed[0].error_code == "parse_failed"
    assert len(service.requests) == 2


def test_task_inputs_keep_the_worksheet_name_from_their_processing_revision(tmp_path):
    from openpyxl import Workbook

    root, identity, _ = question_knowledge(tmp_path)
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "VALUE_17"
    source = tmp_path / "book.xlsx"
    book.save(source)
    responses = [
        {"content": json.dumps({"description": "Value", "content": "VALUE_17"})},
        {"content": json.dumps({"concepts": [], "entities": []})},
    ] * 2
    contexts, results = [], []
    with StreamTransport(responses) as service:
        (root / ".env").write_text(f"LLM_API_KEY=test-key\nOPENAI_API_BASE={service.url}\n")
        for name in ("Alpha", "Renamed Alpha"):
            book.active.title = name
            book.save(source)
            context = ExecutionContext(storage_identity=identity)
            result = import_document(
                root, source, context=context, scope=storage_scope(root, identity)
            )
            assert result.status == "added", result.message
            contexts.append(context)
            results.append(result)
    source.unlink()
    assert results[0].units[0].unit_id == results[1].units[0].unit_id
    for context, name in zip(contexts, ("Alpha", "Renamed Alpha"), strict=True):
        detail = read_knowledge_task(root, context.storage_execution.task_id, limit=128)
        originals = [stage.input for stage in detail.stages if stage.unit_revision_id is not None]
        assert originals and all(original.unit_label == name for original in originals)
    assert len(service.requests) == 4


def test_stage_records_real_blob_costs_and_writer_boundaries_without_double_counting(tmp_path):
    from ultrakb.application.knowledge_tasks import read_knowledge_stage
    from ultrakb.application.task_stages import task_stage
    from ultrakb.storage.blobs import put_blob, read_bytes
    from ultrakb.storage.coordination import storage_stage

    root, identity, _ = question_knowledge(tmp_path)
    context = ExecutionContext(storage_identity=identity)
    payload = b"One retained payload, one deduplicated hit."
    with context.begin(root):
        with task_stage(context, "probe.storage.v1") as progress:
            with storage_stage(root, identity) as db:
                blob = put_blob(db, payload)
                assert put_blob(db, payload) == blob
                assert read_bytes(db, blob.digest) == payload
            pending = read_knowledge_stage(
                root,
                context.storage_execution.task_id,
                progress["stage_attempt_id"],
                identity=identity,
            )
            assert pending.stage.observation_coverage == "pending"
        observed = read_knowledge_stage(
            root, context.storage_execution.task_id, progress["stage_attempt_id"], identity=identity
        )
    assert observed.stage.observation_coverage == "recorded"
    costs = observed.stage.observation.storage
    assert costs.writer_attempts == costs.writer_commits == 2
    assert costs.writer_rollbacks == costs.writer_errors == 0
    assert costs.blob_write_attempts == 1 and costs.blob_write_bytes == len(payload)
    assert costs.blob_read_attempts == 2 and costs.blob_read_bytes == 2 * len(payload)
    assert costs.blob_dedup_hits == 1
    assert costs.writer_hold_seconds >= costs.writer_max_hold_seconds > 0
    assert costs.blob_read_seconds > 0 and costs.blob_write_seconds > 0


@pytest.mark.parametrize("enabled", [False, True])
def test_nested_cost_scopes_keep_core_stages_and_do_not_sum_child_work_twice(tmp_path, enabled):
    from io import BytesIO

    from ultrakb.application.knowledge_tasks import read_knowledge_stage
    from ultrakb.application.task_stages import task_stage
    from ultrakb.storage.blobs import copy_blob, put_blob
    from ultrakb.storage.coordination import storage_stage

    root, identity, _ = question_knowledge(tmp_path)
    context = ExecutionContext(storage_identity=identity, observe_storage_costs=enabled)
    with context.begin(root):
        with task_stage(context, "probe.outer.v1") as outer:
            with storage_stage(root, identity) as db:
                put_blob(db, b"outer-first")
                with task_stage(context, "probe.inner.v1") as inner:
                    blob = put_blob(db, b"inner-payload")
                    copied = BytesIO()
                    copy_blob(db, blob.digest, copied)
                    assert copied.getvalue() == b"inner-payload"
                put_blob(db, b"outer-last")
    details = [
        read_knowledge_stage(
            root, context.storage_execution.task_id, stage["stage_attempt_id"], identity=identity
        )
        for stage in (outer, inner)
    ]
    assert all(detail.stage.state == "completed" for detail in details)
    if enabled:
        outer_cost, inner_cost = (detail.stage.observation.storage for detail in details)
        assert outer_cost.writer_attempts == outer_cost.writer_commits == 2
        assert inner_cost.writer_attempts == inner_cost.writer_commits == 1
        assert outer_cost.blob_write_bytes == len(b"outer-firstouter-last")
        assert inner_cost.blob_write_bytes == inner_cost.blob_read_bytes == len(b"inner-payload")
    else:
        assert all(detail.stage.observation is None for detail in details)
        assert all(detail.stage.observation_coverage == "not_recorded" for detail in details)


@pytest.mark.parametrize("corrupt", [False, True])
def test_private_compile_materialization_records_actual_reads_including_integrity_failures(
    tmp_path, corrupt
):
    from ultrakb.application.knowledge_tasks import read_knowledge_stage
    from ultrakb.application.task_stages import task_stage
    from ultrakb.locks import PrivateMaterialization
    from ultrakb.storage.blobs import blob_path, materialize_private, put_blob
    from ultrakb.storage.coordination import storage_stage

    root, identity, _ = question_knowledge(tmp_path)
    context = ExecutionContext(storage_identity=identity)
    payload = b"Complete Markdown and an attachment read from retained CAS."
    with context.begin(root), storage_stage(root, identity) as db:
        blob = put_blob(db, payload)
        if corrupt:
            payload = b"Truncated CAS payload"
            blob_path(root, blob.digest).write_bytes(payload)
        with PrivateMaterialization() as workspace:
            with task_stage(context, "probe.private.v1") as stage:
                if corrupt:
                    with pytest.raises(ValueError, match="integrity"):
                        materialize_private(db, blob, workspace, "raw/input.md")
                    assert not (workspace.directory / "raw/input.md").exists()
                else:
                    materialize_private(db, blob, workspace, "raw/input.md")
                    assert (workspace.directory / "raw/input.md").read_bytes() == payload
    detail = read_knowledge_stage(
        root, context.storage_execution.task_id, stage["stage_attempt_id"], identity=identity
    )
    costs = detail.stage.observation.storage
    assert costs.blob_read_attempts == 1
    assert costs.blob_read_bytes == len(payload)
    assert costs.blob_read_errors == int(corrupt)
    assert costs.blob_read_seconds > 0
    assert costs.blob_write_attempts == costs.writer_attempts == 0


def test_invalid_optional_observations_do_not_hide_core_task_facts_or_usage(imported_task):
    from ultrakb.application.knowledge_tasks import capture_task_identity, read_knowledge_stage
    from ultrakb.storage.coordination import storage_stage

    root, task_id, _ = imported_task
    identity = capture_task_identity(root)
    before = read_knowledge_task(root, task_id, identity=identity, limit=128)
    stage = next(item for item in before.stages if item.input and item.observation)
    with storage_stage(root, identity) as db, db.transaction() as connection:
        connection.execute(
            "UPDATE ukb_stage_attempts SET observation = ? WHERE stage_attempt_id = ?",
            ('{"schema_version":999}', stage.stage_attempt_id),
        )
    detail = read_knowledge_stage(root, task_id, stage.stage_attempt_id, identity=identity)
    assert detail.stage.state == "completed"
    assert detail.stage.observation is None
    assert detail.stage.observation_coverage == "partial"
    assert detail.stage.observation_unavailable_reason == "invalid_or_unsupported_observation"
    assert detail.current_usage.requests == detail.cumulative_usage.requests == 2
    assert read_knowledge_task(root, task_id).execution.state == "completed"


def test_admission_links_the_temporary_input_to_its_managed_original(imported_task):
    root, task_id, _ = imported_task
    detail = read_knowledge_task(root, task_id, limit=128)
    admission = next(stage for stage in detail.stages if stage.stage_code == "import.admit.v1")
    assert admission.source_revision_id is not None
    assert admission.input is not None
    observed = admission.observation.input
    assert observed.temporary_input_id
    assert observed.input_digest == admission.input.original_digest
    assert observed.input_bytes == admission.input.original_bytes
    assert observed.input_name == admission.input.name
    assert {stage.input.source_id for stage in detail.stages if stage.input} == {
        admission.input.source_id
    }


def test_interrupted_observation_preserves_known_input_and_unknown_final_cost(tmp_path):
    from ultrakb.application.knowledge_tasks import read_knowledge_stage
    from ultrakb.stage_observations import InputObservation, StageObservation
    from ultrakb.storage.execution_store import ExecutionStore

    root, identity, _ = question_knowledge(tmp_path)
    store = ExecutionStore(root, identity)
    execution = store.begin("probe", fingerprint="test", max_calls=0)
    stage_id = store.start_stage(
        execution.execution_id,
        "probe.interrupted.v1",
        observation=StageObservation(
            input=InputObservation(input_name="input.xlsx", input_bytes=17)
        ),
    )
    store.finish_stage(stage_id, state="interrupted")
    store.finish(execution.execution_id, state="interrupted")
    detail = read_knowledge_stage(root, execution.task_id, stage_id, identity=identity)
    assert detail.stage.duration_seconds is None and detail.stage.ended_at is None
    assert detail.stage.observation.input.input_bytes == 17
    assert detail.stage.observation.storage is None
    assert detail.stage.observation_coverage == "partial"
