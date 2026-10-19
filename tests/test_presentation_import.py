"""Hidden slides and speaker notes retain physical page identity at public seams."""

import shutil
from pathlib import Path

import pytest
from native_catalog_fixtures import original_bytes, retained_pdf_bytes

pytest_plugins = ("test_office_import", "block_fixtures", "processing_fixtures")

FIXTURE = Path(__file__).parent / "fixtures/office/slides.pptx"


@pytest.fixture
def frozen_slide_conversion(monkeypatch):
    import io

    import pymupdf
    from pptx import Presentation

    from ultrakb.file_digest import hash_file
    from ultrakb.office.records import OfficeConversion, Slide

    identity = {"fixture": "frozen-physical-slides"}

    def convert(kb_dir, source, output, **kwargs):
        presentation = Presentation(source)
        slides = []
        with pymupdf.open() as pdf:
            for ordinal, slide in enumerate(presentation.slides, 1):
                body = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                page = pdf.new_page()
                if body:
                    page.insert_text((72, 72), body)
                for shape in slide.shapes:
                    if shape.shape_type == 13:
                        page.insert_image((72, 100, 136, 164), stream=shape.image.blob)
                slides.append(
                    Slide(
                        ordinal=ordinal,
                        page=ordinal,
                        name=str(ordinal),
                        hidden=False,
                        body=body,
                        notes="",
                    )
                )
            pdf.save(output)
        artifact = {"url": "fixture://office", "sha256": "a" * 64, "bytes": 1}
        record = OfficeConversion(
            input_digest=hash_file(source),
            version="26.2.6.3",
            build_id="8221e31b3ac356a1623c672912a3d2b492f7e3d1",
            python_version="3.12.14",
            runtime_fingerprint="b" * 64,
            processing_identity=identity,
            probe={},
            archive=artifact,
            source_archive=artifact,
            fonts={},
            licenses={},
            elapsed_seconds=0.0,
            filter="impress_pdf_Export",
            detected_filter="Impress MS PowerPoint 2007 XML",
            diagnostics=[],
            pages=len(slides),
            pdf_digest=hash_file(output),
            pdf_fonts=[],
            font_substitutions=[],
            font_observations=[],
            slides=slides,
        )
        receipt = output.with_suffix(".office.json")
        receipt.write_text(record.model_dump_json())
        return receipt

    monkeypatch.setattr("ultrakb.office.convert.convert_office", convert)
    monkeypatch.setattr("ultrakb.office.runtime.processing_identity", lambda _: identity)

    def original(path, body=""):
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64), False)
        pixmap.clear_with(128)
        slide.shapes.add_picture(io.BytesIO(pixmap.tobytes("png")), Inches(1), Inches(1))
        if body:
            slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text = body
        presentation.save(path)
        return path

    return original


@pytest.fixture
def single_slide_model(pdf_model, monkeypatch, live_llm):
    import litellm

    calls = []

    def completion(**kwargs):
        calls.append(kwargs)
        return live_llm.completion(**kwargs)

    async def acompletion(**kwargs):
        calls.append(kwargs)
        return await live_llm.acompletion(**kwargs)

    monkeypatch.setattr(litellm, "completion", completion)
    monkeypatch.setattr(litellm, "acompletion", acompletion)
    return calls


@pytest.mark.parametrize("limit", [10, 0])
def test_image_only_presentation_retains_pdf_without_publishing_knowledge(
    kb_dir, tmp_path, monkeypatch, frozen_slide_conversion, limit
):
    import pymupdf

    from ultrakb.application.documents import import_document
    from ultrakb.application.settings import apply_kb_config_patch
    from ultrakb.application.settings_data import KbConfigPatchRequest
    from ultrakb.documents import read_document_source

    apply_kb_config_patch(
        kb_dir, KbConfigPatchRequest(kb=str(kb_dir), config={"pdf_short_max_pages": limit})
    )
    path = frozen_slide_conversion(tmp_path / "image-only.pptx")
    original = path.read_bytes()

    def unexpected(**kwargs):
        pytest.fail("Textless presentations cannot request compilation or a navigation tree")

    monkeypatch.setattr("litellm.completion", unexpected)
    monkeypatch.setattr("litellm.acompletion", unexpected)
    result = import_document(
        kb_dir,
        path,
    )
    assert result.status == "blocked", result.message
    assert result.units[0].status == "blocked"
    assert "未生成知识" in result.message
    assert result.units[0].knowledge_revision_id is None
    assert result.model_usage["current"]["requests"] == 0
    unit_id = result.units[0].target_revision_id
    saved = read_document_source(kb_dir, result.source_id, unit_revision_id=unit_id)
    from ultrakb.application.views import published_scope
    from ultrakb.storage.catalog_store import CatalogStore

    catalog = CatalogStore(kb_dir, published_scope(kb_dir).knowledge.identity)
    target = catalog.unit_revision(unit_id)
    processed = catalog.processing(target.processing_id)
    pdf_file = next(file for file in processed.files if file.path.endswith(".pdf"))
    pdf_bytes = catalog.read_processed(processed.processing_id, pdf_file.path)
    assert catalog.read_original(result.source_revision_id) == original == path.read_bytes()
    assert saved["pages"] == 1
    for part in ("body", "notes"):
        selected = read_document_source(
            kb_dir, result.source_id, unit_revision_id=unit_id, pages="1", part=part
        )
        assert selected["units"][0]["parts"][part] == ""
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
        assert len(pdf) == 1 and pdf[0].get_images()
    if limit == 10:
        assert saved["units"][0]["parts"] == {"body": "", "notes": ""}
        assert saved["units"][0]["images"]


@pytest.mark.parametrize("limit", [10, 0])
def test_image_only_replacement_keeps_previous_knowledge_and_retains_new_pdf(
    kb_dir, tmp_path, frozen_slide_conversion, single_slide_model, limit
):
    import pymupdf

    from ultrakb.application.documents import import_document
    from ultrakb.application.settings import apply_kb_config_patch
    from ultrakb.application.settings_data import KbConfigPatchRequest
    from ultrakb.documents import read_document_source

    apply_kb_config_patch(
        kb_dir, KbConfigPatchRequest(kb=str(kb_dir), config={"pdf_short_max_pages": limit})
    )
    path = frozen_slide_conversion(tmp_path / "replacement.pptx", "Original rule.")
    first = import_document(
        kb_dir,
        path,
    )
    assert first.status == "added", first.message
    before = read_document_source(kb_dir, first.source_id)
    calls = len(single_slide_model)
    frozen_slide_conversion(path)
    original = path.read_bytes()
    failed = import_document(
        kb_dir,
        path,
    )
    assert failed.status == "blocked", failed.message
    assert len(single_slide_model) == calls
    assert failed.model_usage["current"]["requests"] == 0
    current = read_document_source(kb_dir, first.source_id)
    for key in (
        "knowledge_revision_id",
        "source_revision_id",
        "content",
        "office",
        "unit_revision_id",
    ):
        assert current[key] == before[key]
    saved = read_document_source(
        kb_dir, first.source_id, unit_revision_id=failed.units[0].target_revision_id
    )
    assert saved["knowledge_revision_id"] is None
    assert original_bytes(kb_dir, saved["source_revision_id"]) == original
    with pymupdf.open(stream=retained_pdf_bytes(kb_dir, saved), filetype="pdf") as pdf:
        assert len(pdf) == 1 and pdf[0].get_images()


def test_recompilation_of_previously_published_textless_slides_does_not_call_model(
    kb_dir, tmp_path, frozen_slide_conversion, single_slide_model, monkeypatch
):
    import asyncio

    from ultrakb.application.documents import import_document
    from ultrakb.application.recompilation import recompile_document, select_recompilation
    from ultrakb.documents import read_document_source

    path = frozen_slide_conversion(tmp_path / "old-textless.pptx")
    # Simulate the prior full-compilation policy to retain an old successful snapshot.
    with monkeypatch.context() as old_policy:
        from dataclasses import replace

        from ultrakb.import_text import preflight_import_text

        def old_preflight(*args, **kwargs):
            return replace(preflight_import_text(*args, **kwargs), unavailable=None)

        old_policy.setattr("ultrakb.import_text.preflight_import_text", old_preflight)
        old_policy.setattr("ultrakb.office.slide_content.require_slide_navigation", lambda _: None)
        first = import_document(
            kb_dir,
            path,
        )
    assert first.status == "added", first.message
    before = read_document_source(kb_dir, first.source_id)
    calls = len(single_slide_model)
    review = select_recompilation(kb_dir, first.source_id, confirmation=True)
    result = asyncio.run(recompile_document(kb_dir, first.source_id, version=review.version))
    assert result.status == "rejected", result.message
    assert "未生成知识" in result.message
    assert len(single_slide_model) == calls
    current = read_document_source(kb_dir, first.source_id)
    assert current["knowledge_revision_id"] == before["knowledge_revision_id"]
    assert current["content"] == before["content"]


@pytest.fixture
def presentation_model(pdf_model, monkeypatch, live_llm):
    import litellm

    prompts = []

    def completion(**kwargs):
        prompt = str(kwargs["messages"])
        prompts.append(prompt)
        return live_llm.completion(**kwargs)

    async def acompletion(**kwargs):
        prompts.append(str(kwargs["messages"]))
        return await live_llm.acompletion(**kwargs)

    monkeypatch.setattr(litellm, "completion", completion)
    monkeypatch.setattr(litellm, "acompletion", acompletion)
    return prompts


@pytest.mark.parametrize("limit", [10, 1])
def test_pptx_retains_hidden_slides_and_notes_without_extra_pages(
    kb_dir, tmp_path, office_runtime, processing_only, limit
):
    from native_processing_inputs import retain_physical_slide_pages
    from native_usage_fixtures import aggregate_usage
    from processing_fixtures import source_processing_context

    from ultrakb.application.settings import apply_kb_config_patch
    from ultrakb.application.settings_data import KbConfigPatchRequest
    from ultrakb.documents import read_document_source

    apply_kb_config_patch(
        kb_dir, KbConfigPatchRequest(kb=str(kb_dir), config={"pdf_short_max_pages": limit})
    )
    path = tmp_path / "演示文稿.pptx"
    shutil.copyfile(FIXTURE, path)
    context = source_processing_context()
    imported = processing_only(path, context=context)
    retain_physical_slide_pages(kb_dir, imported.source_id)
    assert context.executor is not None and context.executor.calls == 0
    assert aggregate_usage(kb_dir)["requests"] == 0
    source = read_document_source(kb_dir, imported.source_id)
    assert source["knowledge_revision_id"] is None
    assert source["length_class"] == ("long" if limit == 1 else "short")
    assert source["execution_mode"] == ("segmented" if limit == 1 else "full")
    assert source["pages"] == 3
    assert "HIDDEN_SLIDE_BODY" in source["content"]
    assert "演讲备注" in source["content"]
    assert "rotate key every 90 days" in source["content"]
    assert [s["hidden"] for s in source["office"]["slides"]] == [False, True, False]
    assert "Alpha visible slide" in source["units"][0]["parts"]["body"]
    assert "HIDDEN_SLIDE_BODY" in source["units"][1]["parts"]["body"]
    page = read_document_source(kb_dir, imported.source_id, pages="3")
    assert page["units"][0]["parts"]["body"].strip() == ""
    assert page["units"][0]["parts"]["notes"] == "Notes-only page evidence."
    assert original_bytes(kb_dir, source["source_revision_id"]) == path.read_bytes()
    notes = read_document_source(kb_dir, imported.source_id, pages="1", part="notes")
    assert "rotate key every 90 days" in notes["content"]
    assert "Alpha visible slide" not in notes["content"]
    assert notes["origin_locators"][0]["part"] == "notes"


def test_pptx_rejects_unusable_summary_capacity_and_retains_source(
    kb_dir, office_runtime, presentation_model
):
    from ultrakb.application.documents import import_document
    from ultrakb.application.settings import apply_kb_config_patch
    from ultrakb.application.settings_data import KbConfigPatchRequest
    from ultrakb.documents import read_document_source

    apply_kb_config_patch(
        kb_dir,
        KbConfigPatchRequest(
            kb=str(kb_dir),
            config={"model": "gpt-4o", "model_capacity": {"max_input_tokens": 1100}},
        ),
    )
    imported = import_document(kb_dir, FIXTURE)
    assert imported.status == "failed", imported.message
    assert "document_summary_capacity_exceeded" in imported.message
    saved = read_document_source(kb_dir, imported.source_id)
    assert saved["knowledge_revision_id"] is None
    assert saved["length_class"] == "short" and saved["execution_mode"] == "segmented"
    processing = saved["target_processing"]
    assert processing["input_limit"] == 1100
    assert processing["estimated_input_tokens"] > processing["input_limit"]
    assert original_bytes(kb_dir, saved["source_revision_id"]) == FIXTURE.read_bytes()
    assert retained_pdf_bytes(kb_dir, saved).startswith(b"%PDF-")
    assert saved["units"][2]["parts"]["notes"] == "Notes-only page evidence."


def test_notes_only_change_rebuilds_index_and_cache_recovers_both_parts(
    kb_dir, tmp_path, office_runtime, presentation_model, wide_short_prompt, live_llm
):
    import json
    import sqlite3

    from pageindex import IndexConfig

    from ultrakb.application.documents import import_document
    from ultrakb.application.settings import apply_kb_config_patch
    from ultrakb.application.settings_data import KbConfigPatchRequest
    from ultrakb.block_package import create_index_client
    from ultrakb.condb_storage import ConDBPageIndexStorage
    from ultrakb.documents import read_document_source
    from ultrakb.index_llm import PageIndexLLM
    from ultrakb.index_location import IndexLocation
    from ultrakb.llm_execution import CompletionExecutor, RoleBindings

    # Force segmentation at the full-request seam while leaving room for the
    # segmented summary protocol; an undersized limit stops before cache checks.
    apply_kb_config_patch(
        kb_dir,
        KbConfigPatchRequest(
            kb=str(kb_dir),
            config={
                "model": "gpt-4o",
                "model_capacity": {"max_input_tokens": 4000},
            },
        ),
    )
    source = tmp_path / "portable.pptx"
    shutil.copyfile(FIXTURE, source)
    imported = import_document(kb_dir, source)
    assert imported.status == "added", imported.message
    saved = read_document_source(kb_dir, imported.source_id)
    assert saved["length_class"] == "short" and saved["execution_mode"] == "segmented"
    processing = saved["target_processing"]
    assert processing["input_limit"] == 4000
    assert processing["estimated_input_tokens"] > processing["input_limit"]
    import base64

    from ultrakb.office.slide_package import SlidePackage

    package = tmp_path / "slides.okpi"
    package.write_text(
        SlidePackage(
            pdf=base64.b64encode(retained_pdf_bytes(kb_dir, saved)).decode("ascii"),
            pdf_digest=saved["office"]["pdf_digest"],
            slides=saved["office"]["slides"],
        ).model_dump_json()
    )
    store = tmp_path / "standalone-index"
    # This regression needs a real navigation decision, not new node summaries
    # or another document description for each copy of the same frozen source.
    index_options = IndexConfig(
        if_add_node_summary=False,
        if_add_doc_description=False,
        llm_client=PageIndexLLM(
            CompletionExecutor(RoleBindings("deepseek/deepseek-flash", live_llm.credentials))
        ),
    )
    with ConDBPageIndexStorage(IndexLocation.package(store)) as database:
        client = create_index_client(
            storage_path=str(store),
            storage=database,
            model="gpt-4o",
            index_config=index_options,
        )
        collection = client.collection()
        tree_prompts = len(presentation_model)
        first = collection.add(str(package))
        assert any(
            "PHYSICAL SLIDE STRUCTURE" in prompt for prompt in presentation_model[tree_prompts:]
        )
        before = collection.get_document(first, include_text=True)
        assert before["page_count"] == 3
        nodes = []

        def collect(items):
            for node in items:
                nodes.append(node)
                collect(node.get("nodes", []))

        collect(before["structure"])
        assert all(node["title_origin"] == "generated" for node in nodes)
        assert any(
            node["anchor"]["part"] == "notes"
            and node["anchor"]["unit"] == 1
            and "rotate key every 90 days" in node["text"]
            for node in nodes
        )
        assert any("AFTER_IMAGE_MARKER" in node["anchor"]["excerpt"] for node in nodes)
        for node in nodes:
            anchor = node["anchor"]
            start, end = anchor["range"]
            # Coordinates address the index's frozen original, before paths
            # are materialized for a different KB/source location.
            indexed_part = before["metadata"]["page_parts"][anchor["unit"] - 1][anchor["part"]]
            assert indexed_part[start:end] == anchor["excerpt"]
        content = json.loads(package.read_text())
        content["slides"][2]["notes"] = "Changed notes, identical PDF bytes."
        package.write_text(json.dumps(content))
        second = collection.add(str(package))
        assert second != first
    moved = tmp_path / "moved-index"
    shutil.copytree(store, moved)
    shutil.rmtree(store)
    source.unlink()
    package.unlink()
    with ConDBPageIndexStorage(IndexLocation.package(moved)) as database:
        collection = create_index_client(
            storage_path=str(moved), storage=database, model="gpt-4o"
        ).collection()
        expected = collection.get_page_content(first, "1-3")
        assert Path(expected[1]["images"][0]["path"]).is_file()
        assert expected[2]["parts"]["notes"] == "Notes-only page evidence."
        before = collection.get_document(first, include_text=True)
        with sqlite3.connect(moved / "context.sqlite") as db:
            cleared = db.execute(
                "UPDATE okb_documents SET payload = json_set(payload, '$.pages', NULL)"
            )
            assert cleared.rowcount == 2
            assert (
                db.execute(
                    "SELECT COUNT(*) FROM okb_documents "
                    "WHERE json_extract(payload, '$.pages') IS NULL"
                ).fetchone()[0]
                == 2
            )
        calls = len(presentation_model)
        assert collection.get_page_content(first, "1-3") == expected
        assert collection.get_document(first, include_text=True) == before
        assert (
            collection.get_page_content(second, "3")[0]["parts"]["notes"]
            == "Changed notes, identical PDF bytes."
        )
        assert len(presentation_model) == calls


def test_presentation_api_and_cli_select_the_same_speaker_notes(
    kb_dir, office_runtime, presentation_model
):
    import json

    from click.testing import CliRunner
    from fastapi.testclient import TestClient

    from ultrakb.api import create_app
    from ultrakb.cli import cli
    from ultrakb.config import register_kb_alias

    register_kb_alias("presentation", kb_dir)
    with TestClient(create_app()) as client:
        response = client.post(
            "/api/v1/add",
            data={"kb": "presentation", "stream": "false"},
            files={"files": ("slides.pptx", FIXTURE.read_bytes())},
        )
        assert response.status_code == 200, response.text
        result = response.json()["files"][0]
        assert result["status"] == "added", result
        read = client.post(
            "/api/v1/document/source",
            json={
                "kb": "presentation",
                "hash": result["source_id"],
                "pages": "2",
                "part": "notes",
            },
        )
    assert read.status_code == 200, read.text
    notes = read.json()
    assert notes["part"] == "notes" and notes["pages"] == 3
    assert "Hidden speaker note" in notes["content"] and "HIDDEN_SLIDE_BODY" not in notes["content"]
    output = CliRunner().invoke(
        cli,
        ["--kb-dir", str(kb_dir), "source", result["source_id"], "--pages", "2", "--part", "notes"],
    )
    assert output.exit_code == 0, output.output
    assert json.loads(output.output)["content"] == notes["content"]


def test_note_anchor_cannot_claim_body_coordinates():
    from pageindex.index.page_parts_policy import PagePartsPolicy

    policy = PagePartsPolicy(
        {
            "unit_count": 1,
            "page_parts": [
                {
                    "body": "Alpha visible slide",
                    "notes": "Speaker note only: rotate key every 90 days.",
                }
            ],
        }
    )
    entry = {
        "structure": "1",
        "title": "Key rotation",
        "title_origin": "generated",
        "physical_index": 1,
        "anchor": {"unit": 1, "part": "body", "range": [0, 17], "excerpt": "Speaker note only"},
    }
    policy.normalize_anchor(entry)
    assert not policy.valid(entry)
    with pytest.raises(ValueError, match="Unverified physical slide anchors"):
        policy.validate_order([entry])
    assert policy.parts[0]["notes"] == "Speaker note only: rotate key every 90 days."
