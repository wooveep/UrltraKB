"""Small real workbook inputs, independent of test-module plugin registration."""

import pytest


@pytest.fixture
def three_sheets(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA_SHEET"
    book.create_sheet("Beta")["A1"] = "BETA_SHEET"
    book.create_sheet("Gamma")["A1"] = "GAMMA_SHEET"
    book["Gamma"].sheet_state = "hidden"
    path = tmp_path / "three.xlsx"
    book.save(path)
    return path
