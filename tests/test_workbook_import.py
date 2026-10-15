"""A workbook publishes each worksheet independently at public application boundaries."""

from collections import Counter

pytest_plugins = ("test_office_import", "block_fixtures", "workbook_fixtures")


def worksheet_prompt_markers(messages, markers=("ALPHA_SHEET", "BETA_SHEET", "GAMMA_SHEET")):
    """Observe the current source, excluding earlier publications in wiki context."""
    text = []
    for message in messages:
        content = message.get("content")
        if isinstance(content, list):
            content = "\n".join(
                block.get("text", "") for block in content if isinstance(block, dict)
            )
        if isinstance(content, str):
            text.append(content)
    original = next(
        (
            content
            for content in text
            if content.startswith(("New document: ", "Original document: "))
        ),
        "\n".join(text),
    )
    return [marker for marker in markers if marker in original]


def test_one_sheet_failure_keeps_other_publications_and_retry_skips_them(
    kb_dir, three_sheets, pdf_model, monkeypatch
):
    import litellm

    from ultrakb.application.documents import import_document
    from ultrakb.application.sources import source_inventory
    from ultrakb.documents import read_document_source

    original = litellm.completion
    calls = Counter()
    fail = True

    def completion(**kwargs):
        for sheet in worksheet_prompt_markers(kwargs["messages"]):
            calls[sheet] += 1
            if fail and sheet == "BETA_SHEET":
                raise ValueError("Local transport rejects only Beta")
        return original(**kwargs)

    async def acompletion(**kwargs):
        return completion(**kwargs)

    monkeypatch.setattr(litellm, "completion", completion)
    monkeypatch.setattr(litellm, "acompletion", acompletion)
    first = import_document(kb_dir, three_sheets)
    assert first.status == "partial", first.message
    units = {unit.name: unit for unit in first.units}
    assert {name: item.status for name, item in units.items()} == {
        "Alpha": "completed",
        "Beta": "failed",
        "Gamma": "completed",
    }
    retained = read_document_source(kb_dir, first.source_id, unit_id=units["Gamma"].unit_id)
    assert retained["sheet"]["state"] == "hidden" and "GAMMA_SHEET" in retained["content"]
    before = calls.copy()
    fail = False
    resumed = import_document(kb_dir, three_sheets)
    assert resumed.status == "added", resumed.message
    assert calls["BETA_SHEET"] > before["BETA_SHEET"]
    for name in ("Alpha", "Gamma"):
        saved = next(item for item in resumed.units if item.name == name)
        assert saved.knowledge_revision_id == units[name].knowledge_revision_id
        assert saved.model_usage["current"]["requests"] == 0
    sources = source_inventory(kb_dir)
    assert len(sources) == 1 and len(sources[0]["units"]) == 3


def test_segmented_sheet_keeps_cell_locations_and_recompiles_offline(
    kb_dir, three_sheets, pdf_model, monkeypatch, wide_short_prompt
):
    import asyncio

    import litellm
    from openpyxl import load_workbook

    from ultrakb.application.documents import import_document
    from ultrakb.application.recompilation import recompile_document, select_recompilation
    from ultrakb.application.settings import apply_kb_config_patch
    from ultrakb.application.settings_data import KbConfigPatchRequest
    from ultrakb.documents import read_document_source

    # Only Gamma's real segmented generation and cached offline recompilation
    # are under test. XLSX requires a visible sheet beside the hidden target.
    book = load_workbook(three_sheets)
    book["Alpha"]["A1"] = None
    book.remove(book["Beta"])
    book.save(three_sheets)
    book.close()
    apply_kb_config_patch(
        kb_dir,
        KbConfigPatchRequest(
            # Planning and cell provenance fit; the wide short prompt selects
            # segmented execution without making the shared schema unusable.
            kb=str(kb_dir),
            config={"model_capacity": {"max_input_tokens": 4000, "tokenizer_model": "gpt-4o-mini"}},
        ),
    )
    original, async_original = litellm.completion, litellm.acompletion
    structures = []

    def completion(**kwargs):
        prompt = str(kwargs["messages"])
        if "CONTENT BLOCK STRUCTURE" not in prompt:
            return original(**kwargs)
        structures.append(prompt)
        return original(**kwargs)

    async def acompletion(**kwargs):
        if "CONTENT BLOCK STRUCTURE" in str(kwargs["messages"]):
            structures.append(str(kwargs["messages"]))
        return await async_original(**kwargs)

    monkeypatch.setattr(litellm, "completion", completion)
    monkeypatch.setattr(litellm, "acompletion", acompletion)
    imported = import_document(kb_dir, three_sheets)
    assert imported.status == "added", imported.message
    units = {unit.name: unit for unit in imported.units}
    assert {name: unit.status for name, unit in units.items()} == {
        "Alpha": "empty",
        "Gamma": "completed",
    }
    assert units["Alpha"].model_usage["current"]["requests"] == 0
    gamma = next(unit for unit in imported.units if unit.name == "Gamma")
    saved = read_document_source(kb_dir, imported.source_id, unit_id=gamma.unit_id, blocks="1")
    assert saved["execution_mode"] == "segmented" and saved["length_class"] == "short"
    assert saved["sheet"]["name"] == "Gamma" and saved["sheet"]["state"] == "hidden"
    assert saved["content"] and structures
    assert any(
        item.get("sheet_cell", {}).get("cell", {}).get("coordinate") == "A1"
        for item in saved["origin_locators"]
    )
    before = len(structures)
    three_sheets.unlink()
    selected = select_recompilation(
        kb_dir, imported.source_id, confirmation=True, unit_id=gamma.unit_id
    )
    result = asyncio.run(
        recompile_document(
            kb_dir, imported.source_id, unit_id=gamma.unit_id, version=selected.version
        )
    )
    assert result.status == "compiled", result.message
    assert len(structures) == before
    again = read_document_source(kb_dir, imported.source_id, unit_id=gamma.unit_id, blocks="1")
    assert again["content"] == saved["content"]


def test_sparse_sheet_preserves_types_formulas_hidden_cells_and_disjoint_ranges(kb_dir, tmp_path):
    import json
    from datetime import date
    from zipfile import ZIP_DEFLATED, ZipFile

    from office_publication_fixtures import publish_workbook
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from ultrakb.documents import read_document_source

    # Typed-cell and pinned-reader assertions consume the real retained worksheet.
    path = tmp_path / "sparse.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Typed values"
    sheet["A1"] = "0012"
    sheet["B2"] = 42
    sheet["B2"].number_format = "00000"
    sheet["D2"] = date(2025, 1, 2)
    sheet["E2"] = "=B2+1"
    sheet["F2"] = "=B2+2"
    sheet["G2"] = True
    sheet["XFD100000"] = "TAIL_VALUE"
    sheet["XFD1048576"].font = Font(bold=True)
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True
    sheet.merge_cells("H3:J4")
    book.save(path)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b'<c r="E2"><f>B2+1</f><v></v></c>', b'<c r="E2"><f>B2+1</f><v>43</v></c>'
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    imported = publish_workbook(kb_dir, path)
    assert imported.status == "added", imported.message
    saved = read_document_source(kb_dir, imported.source_id, unit_id=imported.units[0].unit_id)
    cells = {
        item["sheet_cell"]["cell"]["coordinate"]: item["sheet_cell"]["cell"]
        for item in saved["origin_locators"]
        if item.get("sheet_cell")
    }
    assert len(cells) == 7 and saved["characters"] < 2000
    assert cells["A1"]["value"] == "0012" and cells["B2"]["display"] == "00042"
    assert cells["D2"]["value"] == "2025-01-02T00:00:00" and cells["D2"]["data_type"] == "d"
    assert cells["E2"]["formula"] == "=B2+1" and cells["E2"]["cached"] == 43
    assert cells["F2"]["cache_status"] == "missing" and cells["F2"]["cached"] is None
    assert cells["G2"]["value"] is True
    assert cells["B2"]["hidden_row"] and cells["B2"]["hidden_column"]
    assert saved["sheet"]["merged_ranges"] == ["H3:J4"]
    selected = read_document_source(
        kb_dir, imported.source_id, unit_id=imported.units[0].unit_id, cells="A1:B2,XFD100000"
    )
    assert "0012" in selected["content"] and "TAIL_VALUE" in selected["content"]
    assert "=B2+1" not in selected["content"]
    # The query tool uses the same retained source through its pinned reader.
    import asyncio

    from agents import Agent
    from test_answer_evidence_regression import invoke

    from ultrakb.agent.query_evidence import restrict_query_agent
    from ultrakb.application.query_views import resolve_query_views

    selection = resolve_query_views(kb_dir, "Read worksheet values")
    view = selection.views[0]
    name = next(p for p in view.files if p.endswith(".content.json"))
    agent = restrict_query_agent(Agent(name="sheet-reader"), selection)
    result = asyncio.run(
        invoke(
            agent,
            "get_cell_content",
            doc_name=name.removeprefix("sources/").removesuffix(".content.json"),
            cells="A1:B2,E2:F2,XFD100000",
        )
    )
    payload, _ = json.JSONDecoder().raw_decode(result[result.index("{") :])
    projected = {cell["coordinate"]: cell for cell in payload["cells"]}
    assert projected["E2"]["formula"] == "=B2+1" and projected["E2"]["cached"] == 43
    assert projected["F2"]["cache_status"] == "missing"
    assert projected["B2"]["hidden_row"] and projected["B2"]["hidden_column"]
    assert "TAIL_VALUE" in payload["content"] and "00042" in payload["content"]
    assert imported.source_revision_id in result
    assert [region["char_range"] for region in payload["source_context"]["regions"]] == (
        payload["source_spans"][:8]
    )
    assert "read_id" not in payload
    assert (
        read_document_source(kb_dir, imported.source_id, unit_id=imported.units[0].unit_id) == saved
    )
    frozen_path = view.scope.wiki_dir / name
    original_bytes = frozen_path.read_bytes()
    try:
        frozen_path.write_bytes(b"{}")
        assert "changed after selection" in asyncio.run(
            invoke(
                agent,
                "get_cell_content",
                doc_name=name.removeprefix("sources/").removesuffix(".content.json"),
                cells="A1",
            )
        )
    finally:
        frozen_path.write_bytes(original_bytes)
    assert len(selected["source_spans"]) == 3


def test_workbook_upload_cli_and_single_sheet_recompilation(
    kb_dir, three_sheets, pdf_model, monkeypatch, live_llm
):
    import asyncio
    import json
    from contextlib import contextmanager

    from click.testing import CliRunner
    from fastapi.testclient import TestClient
    from office_publication_fixtures import publish_workbook

    from ultrakb.api import create_app
    from ultrakb.api_helpers import published_input
    from ultrakb.application.recompilation import recompile_document, select_recompilation
    from ultrakb.application.sources import source_inventory
    from ultrakb.cli import cli
    from ultrakb.config import register_kb_alias, resolve_credential_bundle
    from ultrakb.source_catalog import read_source_revision

    register_kb_alias("sheets", kb_dir)
    prepared = []
    uploaded_bytes = three_sheets.read_bytes()

    @contextmanager
    def arrange_other_sheets(*args, **kwargs):
        # Keep actual upload ownership/admission and its three-sheet identity.
        # Other-sheet history is domain state; Gamma uses the real compiler in
        # both the API upload and the selected-sheet recompilation below.
        with published_input(*args, **kwargs) as owned:
            call_start = len(live_llm.calls)
            prior = publish_workbook(
                kb_dir,
                owned.path,
                bundle=resolve_credential_bundle(kb_dir),
                sheet_names=("Alpha", "Beta"),
                record_model_task=False,
            )
            assert prior.status == "added", prior.message
            assert {unit.name for unit in prior.units} == {"Alpha", "Beta"}
            assert all(unit.status == "completed" for unit in prior.units)
            assert len(live_llm.calls) == call_start
            prepared.append(prior)
            yield owned

    monkeypatch.setattr("ultrakb.api_helpers.published_input", arrange_other_sheets)
    upload_call_start = len(live_llm.calls)
    with TestClient(create_app()) as client:
        result = client.post(
            "/api/v1/add",
            data={"kb": "sheets", "stream": "false"},
            files={"files": ("three.xlsx", uploaded_bytes)},
        ).json()["files"][0]
        assert result["status"] == "added", result
        assert result["source_id"] == prepared[0].source_id
        assert result["source_revision_id"] == prepared[0].source_revision_id
        frozen = read_source_revision(kb_dir, result["source_revision_id"])
        assert (kb_dir / frozen.original).read_bytes() == uploaded_bytes
        assert {item["name"] for item in result["units"]} == {"Alpha", "Beta", "Gamma"}
        gamma = next(item for item in result["units"] if item["name"] == "Gamma")
        assert gamma["status"] == "completed"
        assert gamma["model_usage"]["current"]["requests"] > 0
        prior_ids = {unit.unit_id: unit.knowledge_revision_id for unit in prepared[0].units}
        for unit in result["units"]:
            if unit["name"] != "Gamma":
                assert unit["knowledge_revision_id"] == prior_ids[unit["unit_id"]]
                assert unit["model_usage"]["current"]["requests"] == 0
        response = client.post(
            "/api/v1/document/source",
            json={
                "kb": "sheets",
                "hash": result["source_id"],
                "unit_id": gamma["unit_id"],
                "cells": "A1",
            },
        )
    assert response.status_code == 200, response.text
    assert response.json()["sheet"]["name"] == "Gamma"
    assert response.json()["sheet"]["state"] == "hidden"
    assert "GAMMA_SHEET" in response.json()["content"]
    assert 1 <= len(live_llm.calls) - upload_call_start <= 16
    output = CliRunner().invoke(
        cli,
        [
            "--kb-dir",
            str(kb_dir),
            "source",
            result["source_id"],
            "--unit",
            gamma["unit_id"],
            "--cells",
            "A1",
        ],
    )
    assert output.exit_code == 0, output.output
    assert json.loads(output.output)["content"] == response.json()["content"]
    before = {
        unit["unit_id"]: unit["knowledge_revision_id"]
        for unit in source_inventory(kb_dir)[0]["units"]
    }
    selected = select_recompilation(kb_dir, result["source_id"], confirmation=True)
    recompile_call_start = len(live_llm.calls)
    again = asyncio.run(
        recompile_document(
            kb_dir, result["source_id"], unit_id=gamma["unit_id"], version=selected.version
        )
    )
    assert again.status == "compiled", again.message
    recompile_calls = live_llm.calls[recompile_call_start:]
    assert 1 <= len(recompile_calls) <= 16
    assert all(
        worksheet_prompt_markers(call["messages"]) == ["GAMMA_SHEET"] for call in recompile_calls
    )
    assert again.model_usage["current"]["requests"] > 0
    after = {
        unit["unit_id"]: unit["knowledge_revision_id"]
        for unit in source_inventory(kb_dir)[0]["units"]
    }
    assert after.pop(gamma["unit_id"]) != before.pop(gamma["unit_id"])
    assert after == before


def test_retry_one_failed_sheet_uses_retained_input_after_original_is_removed(
    kb_dir, three_sheets, pdf_model, monkeypatch
):
    import litellm

    from ultrakb.application.documents import import_document
    from ultrakb.application.workbook_actions import retry_worksheet

    original = litellm.completion
    failed = {"BETA_SHEET", "GAMMA_SHEET"}
    calls = Counter()

    def completion(**kwargs):
        for name in worksheet_prompt_markers(kwargs["messages"]):
            calls[name] += 1
            if name in failed:
                raise ValueError("Local transport rejects this sheet")
        return original(**kwargs)

    async def acompletion(**kwargs):
        return completion(**kwargs)

    monkeypatch.setattr(litellm, "completion", completion)
    monkeypatch.setattr(litellm, "acompletion", acompletion)
    first = import_document(kb_dir, three_sheets)
    assert first.status == "partial", first.message
    beta = next(unit for unit in first.units if unit.name == "Beta")
    three_sheets.unlink()
    before = calls.copy()
    failed.remove("BETA_SHEET")
    result = retry_worksheet(kb_dir, first.source_id, beta.unit_id)
    assert result.status == "added" and len(result.units) == 1
    assert result.units[0].target_revision_id == beta.target_revision_id
    assert calls["BETA_SHEET"] > before["BETA_SHEET"]
    assert calls["ALPHA_SHEET"] == before["ALPHA_SHEET"]
    assert calls["GAMMA_SHEET"] == before["GAMMA_SHEET"]


def test_historical_workbook_lists_only_its_published_sheet_names(kb_dir, three_sheets):
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook

    from ultrakb.application.query_views import resolve_query_views
    from ultrakb.application.sources import source_inventory
    from ultrakb.documents import read_document_source

    first = publish_workbook(kb_dir, three_sheets)
    pinned = next(
        view.scope
        for view in resolve_query_views(kb_dir, "Read all worksheets").views
        if view.view_id == first.units[0].view_id
    )
    book = load_workbook(three_sheets)
    book["Alpha"].title = "Renamed"
    book.create_sheet("Delta")["A1"] = "NEW_SHEET"
    book.save(three_sheets)
    updated = publish_workbook(kb_dir, three_sheets)
    assert updated.status == "added", updated.message
    historical = source_inventory(kb_dir, scope=pinned)[0]
    assert {unit["name"] for unit in historical["units"]} == {"Alpha", "Beta", "Gamma"}
    alpha = next(unit for unit in first.units if unit.name == "Alpha")
    source = read_document_source(kb_dir, first.source_id, unit_id=alpha.unit_id, scope=pinned)
    assert {unit["name"] for unit in source["available_units"]} == {"Alpha", "Beta", "Gamma"}
    assert source["sheet"]["name"] == "Alpha"
