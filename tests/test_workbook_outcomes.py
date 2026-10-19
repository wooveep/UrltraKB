"""Durable workbook-level and per-unit failures are visible through public adapters."""

import pytest

pytest_plugins = ("test_office_import", "block_fixtures", "workbook_fixtures")


@pytest.mark.parametrize("empty_inventory", [False, True])
def test_inventory_failure_survives_import_and_keeps_old_body(
    kb_dir, three_sheets, empty_inventory
):
    from office_publication_fixtures import publish_workbook

    from ultrakb.application.documents import import_document
    from ultrakb.application.sources import source_inventory
    from ultrakb.documents import read_document_source

    first = publish_workbook(kb_dir, three_sheets)
    assert first.status == "added", first.message
    if empty_inventory:
        import re
        from zipfile import ZipFile

        with ZipFile(three_sheets) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        parts["xl/workbook.xml"] = re.sub(
            b"<sheets>.*?</sheets>", b"<sheets/>", parts["xl/workbook.xml"]
        )
        with ZipFile(three_sheets, "w") as archive:
            for name, value in parts.items():
                archive.writestr(name, value)
    else:
        three_sheets.write_bytes(b"invalid workbook")
    failed = import_document(
        kb_dir,
        three_sheets,
    )
    assert failed.status == "failed"
    read = read_document_source(kb_dir, first.source_id, unit_id=first.units[0].unit_id)
    assert read["status"] == "completed" and "ALPHA_SHEET" in read["content"]
    assert read["source_revision_id"] == first.source_revision_id
    assert read["target_source_revision_id"] == failed.source_revision_id
    inventory = source_inventory(kb_dir)[0]
    assert inventory["status"] == "failed"
    assert "inventory" in inventory["message"].lower()
    assert all(unit["status"] == "completed" for unit in source_inventory(kb_dir)[0]["units"])
    explicit = read_document_source(
        kb_dir,
        first.source_id,
        unit_id=first.units[0].unit_id,
        source_revision_id=failed.source_revision_id,
    )
    assert explicit is None
    from ultrakb.application.views import published_scope
    from ultrakb.storage.catalog_store import CatalogStore

    catalog = CatalogStore(kb_dir, published_scope(kb_dir).knowledge.identity)
    assert catalog.read_original(failed.source_revision_id) == three_sheets.read_bytes()
    old = read_document_source(
        kb_dir,
        first.source_id,
        unit_id=first.units[0].unit_id,
        source_revision_id=first.source_revision_id,
    )
    assert old["status"] == "completed" and old.get("message") is None


def test_invalid_xlsx_cell_coordinate_never_publishes(kb_dir, tmp_path):
    from openpyxl import Workbook

    from ultrakb.application.documents import import_document

    workbook = Workbook()
    workbook.active["XFE1"] = "Outside the XLSX column limit"
    path = tmp_path / "invalid.xlsx"
    workbook.save(path)
    result = import_document(kb_dir, path)
    assert result.status == "failed" and result.units[0].successful_revision_id is None


def test_partial_recompile_has_per_sheet_outcomes_and_is_not_counted_skipped(
    kb_dir, three_sheets, pdf_model, monkeypatch
):
    import asyncio

    import litellm
    from click.testing import CliRunner
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook
    from test_workbook_import import worksheet_prompt_markers

    from ultrakb.api_recompile import iter_recompile
    from ultrakb.cli import cli

    # One completed sheet beside the failed one proves per-unit aggregation and
    # both adapters' PARTIAL display. Three-sheet identity/retry has its own test.
    book = load_workbook(three_sheets)
    book.remove(book["Gamma"])
    book.save(three_sheets)
    book.close()
    imported = publish_workbook(kb_dir, three_sheets)
    assert imported.status == "added", imported.message
    original = litellm.completion

    def complete(**kwargs):
        if "BETA_SHEET" in worksheet_prompt_markers(kwargs["messages"]):
            raise ValueError("Local transport rejects Beta recompile")
        return original(**kwargs)

    monkeypatch.setattr(litellm, "completion", complete)

    async def run():
        return [event async for event in iter_recompile(kb_dir, imported.source_id)]

    final = asyncio.run(run())[-1]
    assert final.get("partial") == 1 and final["skipped"] == 0
    assert {item["status"] for item in final["docs"][0]["units"]} == {"completed", "failed"}
    assert {item["name"]: item["status"] for item in final["docs"][0]["units"]} == {
        "Alpha": "completed",
        "Beta": "failed",
    }
    cli_result = CliRunner().invoke(cli, ["--kb-dir", str(kb_dir), "recompile", imported.source_id])
    assert cli_result.exit_code == 0, cli_result.output
    assert "[PARTIAL]" in cli_result.output and "Beta" in cli_result.output
