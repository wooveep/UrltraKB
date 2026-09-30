"""Read typed cached XLS cells and explicitly mark unavailable formula expressions."""

import hashlib

import xlrd
from openpyxl.utils.cell import get_column_letter

from openkb.workbooks.records import SheetCell, SheetSnapshot
from openkb.workbooks.values import display_value
from openkb.workbooks.xls_records import open_observed_book


def read_xls(path) -> tuple[SheetSnapshot, ...]:
    book = open_observed_book(path)
    try:
        names = [name for name, *_ in book.inventory]
        if len(set(names)) != len(names) or len(book.sheet_names()) != book.nsheets:
            raise ValueError("XLS workbook inventory is inconsistent")
        result = []
        for ordinal, (name, visibility, kind, index) in enumerate(book.inventory, 1):
            basis = {
                "key": "xls:" + hashlib.sha256(name.encode()).hexdigest()[:24],
                "native_id": None,
                "name": name,
                "ordinal": ordinal,
                "state": {0: "visible", 1: "hidden", 2: "veryHidden"}[visibility],
            }
            try:
                if kind != 0 or index < 0:
                    raise ValueError(f"Non-cell XLS sheet type {kind}; no extracted cell evidence")
                sheet = book.sheet_by_index(index)
                result.append(_read_sheet(book, sheet, index, basis))
            except Exception as exc:
                result.append(SheetSnapshot(**basis, error=f"{type(exc).__name__}: {exc}"))
        return tuple(result)
    finally:
        book.release_resources()


def _read_sheet(book, sheet, index, basis):
    formulas = book.formulas.get(index, set())
    hidden_rows = tuple(row + 1 for row, info in sorted(sheet.rowinfo_map.items()) if info.hidden)
    hidden_columns = tuple(
        (column + 1, column + 1)
        for column, info in sorted(sheet.colinfo_map.items())
        if info.hidden
    )
    cells = []
    for row in range(sheet.nrows):
        for column in range(sheet.row_len(row)):
            cell = sheet.cell(row, column)
            formula = (row, column) in formulas
            if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK} and not formula:
                continue
            actual, data_type = _value(book, cell)
            number_format = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
            display = display_value(actual, number_format)
            cells.append(
                SheetCell(
                    coordinate=f"{get_column_letter(column + 1)}{row + 1}",
                    row=row + 1,
                    column=column + 1,
                    value=actual,
                    data_type=data_type,
                    number_format=number_format,
                    display=display,
                    formula_status="unavailable" if formula else "not_formula",
                    cached=actual if formula else None,
                    cache_status=("available" if actual is not None else "missing")
                    if formula
                    else "not_formula",
                    hidden_row=row + 1 in hidden_rows,
                    hidden_column=column in sheet.colinfo_map
                    and bool(sheet.colinfo_map[column].hidden),
                )
            )
    diagnostics = (
        (
            "XLS formula expressions unavailable; BIFF formula cells and stored cached values "
            "are retained. No calculation was performed.",
        )
        if formulas
        else ()
    )
    return SheetSnapshot(
        **basis,
        cells=tuple(cells),
        hidden_rows=hidden_rows,
        hidden_columns=hidden_columns,
        has_objects=index in book.objects,
        merged_ranges=tuple(
            f"{get_column_letter(c0 + 1)}{r0 + 1}:{get_column_letter(c1)}{r1}"
            for r0, r1, c0, c1 in sheet.merged_cells
        ),
        diagnostics=diagnostics,
    )


def _value(book, cell):
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).isoformat(), "d"
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value), "b"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code[cell.value], "e"
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return None, "empty"
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        value = cell.value
        return (int(value) if value.is_integer() else value), "n"
    if cell.ctype == xlrd.XL_CELL_TEXT:
        return cell.value, "s"
    raise ValueError(f"Unsupported XLS cell type {cell.ctype}")
