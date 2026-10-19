"""Real worksheet imports retain independent outcomes in the shared main store."""

import asyncio
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from stream_transport import StreamTransport
from test_storage_questions import question_knowledge

from ultrakb.application.documents import import_document
from ultrakb.application.execution import ExecutionContext
from ultrakb.application.pages import read_page
from ultrakb.application.recompilation import recompile_document
from ultrakb.application.source_retry import retry_source
from ultrakb.application.views import storage_scope
from ultrakb.storage.catalog_store import CatalogStore

pytestmark = pytest.mark.integration


def test_partial_workbook_keeps_other_sheets_and_retry_uses_the_managed_inventory(tmp_path):
    from openpyxl import Workbook

    root, identity, _ = question_knowledge(tmp_path)
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA_SHEET"
    book.create_sheet("Beta")["A1"] = "BETA_SHEET"
    book.create_sheet("Gamma")["A1"] = "GAMMA_SHEET"
    book.create_sheet("Empty")
    source = tmp_path / "book.xlsx"
    book.save(source)
    with ZipFile(source) as archive:
        entries = [(entry, archive.read(entry)) for entry in archive.infolist()]
    with ZipFile(source, "w") as archive:
        for entry, data in entries:
            archive.writestr(
                entry, b"<worksheet broken=" if entry.filename.endswith("sheet2.xml") else data
            )
    original = source.read_bytes()
    responses = [
        item
        for text in ("ALPHA_SHEET", "GAMMA_SHEET")
        for item in (
            {"content": json.dumps({"description": text, "content": text})},
            {"content": json.dumps({"concepts": [], "entities": []})},
        )
    ]
    context = ExecutionContext(storage_identity=identity)
    with StreamTransport(responses) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        result = import_document(root, source, context=context, scope=storage_scope(root, identity))
        assert result.status == "partial", result.message
        units = {unit.name: unit for unit in result.units}
        assert set(units) == {"Alpha", "Beta", "Gamma", "Empty"}
        assert units["Alpha"].status == units["Gamma"].status == "completed"
        assert units["Beta"].status == "failed"
        assert units["Empty"].status == "empty"
        assert units["Beta"].target_revision_id != units["Alpha"].target_revision_id
        assert units["Alpha"].model_usage["current"]["actual_sends"] == 2
        assert units["Gamma"].model_usage["current"]["actual_sends"] == 2
        assert units["Gamma"].model_usage["cumulative"]["actual_sends"] == 2
        assert units["Empty"].model_usage["current"]["actual_sends"] == 0
        assert units["Beta"].model_usage["current"]["actual_sends"] == 0
        public = storage_scope(root, identity)
        assert {unit.unit_id for unit in public.knowledge.input_units()} >= {
            units["Alpha"].unit_id,
            units["Gamma"].unit_id,
        }
        native = read_page(root, f"sources/{units['Alpha'].doc_name}", scope=public)
        assert "ALPHA_SHEET" in native.body
        assert "BETA_SHEET" not in native.body
        rejected = asyncio.run(
            recompile_document(root, result.source_id, unit_id=units["Empty"].unit_id, scope=public)
        )
        assert rejected.status == "blocked"
        source.unlink()
        assert CatalogStore(root, identity).read_original(result.source_revision_id) == original
        resumed = retry_source(root, result.source_id, scope=public)
        assert resumed.status == "partial", resumed.message
        assert {unit.name: unit.target_revision_id for unit in resumed.units} == {
            unit.name: unit.target_revision_id for unit in result.units
        }
        assert resumed.model_usage["cumulative"]["actual_sends"] == 4
        assert len(service.requests) == 4
    assert not (root / ".ultrakb/catalog").exists()
    assert not (root / ".ultrakb/hashes.json").exists()


def test_binary_workbook_uses_the_same_managed_unit_pipeline(tmp_path):
    import yaml

    root, identity, _ = question_knowledge(tmp_path)
    config_path = root / ".ultrakb/config.yaml"
    config = yaml.safe_load(config_path.read_text())
    config["model_policy"]["max_calls"] = 8
    config_path.write_text(yaml.safe_dump(config))
    source = tmp_path / "book.xls"
    source.write_bytes(Path("tests/fixtures/office/typed-sheets.xls").read_bytes())
    responses = [
        item
        for text in ("Alpha", "Beta", "Gamma")
        for item in (
            {"content": json.dumps({"description": text, "content": text})},
            {"content": json.dumps({"concepts": [], "entities": []})},
        )
    ]
    with StreamTransport(responses) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        result = import_document(root, source, scope=storage_scope(root, identity))
        assert result.status == "added", result.message
        assert {unit.name for unit in result.units} == {"Alpha", "Beta", "Gamma"}
        assert all(unit.status == "completed" for unit in result.units)
        source.unlink()
        repeated = retry_source(root, result.source_id, scope=storage_scope(root, identity))
        assert repeated.status == "skipped", repeated.message
        assert repeated.model_usage["current"]["actual_sends"] == 0
        assert repeated.model_usage["cumulative"]["actual_sends"] == 6
        assert len(service.requests) == 6
        from ultrakb.application.storage_maintenance import backup_knowledge, restore_knowledge

        backup = tmp_path / "backup"
        backup_knowledge(root, backup, scope=storage_scope(root, identity))
        destination = tmp_path / "restored"
        restored = restore_knowledge(backup, destination)
        restored_scope = storage_scope(destination, restored.identity)
        recovered = retry_source(destination, result.source_id, scope=restored_scope)
        assert recovered.status == "skipped", recovered.message
        assert {unit.unit_id for unit in recovered.units} == {unit.unit_id for unit in result.units}
        assert recovered.model_usage["cumulative"]["actual_sends"] == 6
        assert len(service.requests) == 6
    assert not (root / ".ultrakb/catalog").exists()
    assert not (root / ".ultrakb/hashes.json").exists()


def test_workbook_reordering_keeps_unit_identity_and_old_publications_keep_old_cells(tmp_path):
    from openpyxl import Workbook

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "OLD_ALPHA"
    book.create_sheet("Beta")["A1"] = "BETA"
    book.save(source)
    responses = [
        item
        for text in ("OLD_ALPHA", "BETA", "BETA", "NEW_ALPHA")
        for item in (
            {"content": json.dumps({"description": text, "content": text})},
            {"content": json.dumps({"concepts": [], "entities": []})},
        )
    ]
    with StreamTransport(responses) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        assert first.status == "added", first.message
        old = storage_scope(root, identity)
        old_units = {unit.name: unit for unit in first.units}
        book.move_sheet("Beta", -1)
        book["Alpha"]["A1"] = "NEW_ALPHA"
        book.save(source)
        second = import_document(root, source, scope=old)
        assert second.status == "added", second.message
        new_units = {unit.name: unit for unit in second.units}
        assert {name: unit.unit_id for name, unit in new_units.items()} == {
            name: unit.unit_id for name, unit in old_units.items()
        }
        assert second.source_revision_id != first.source_revision_id
        page = f"sources/{old_units['Alpha'].doc_name}"
        assert "OLD_ALPHA" in read_page(root, page, scope=old).body
        assert "NEW_ALPHA" in read_page(root, page, scope=storage_scope(root, identity)).body
        assert len(service.requests) == 8


def test_failed_workbook_inventory_does_not_break_the_last_complete_sheet_identities(tmp_path):
    from openpyxl import Workbook

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA"
    book.create_sheet("Beta")["A1"] = "BETA"
    book.save(source)
    original = source.read_bytes()
    responses = [
        item
        for text in ("ALPHA", "BETA", "ALPHA", "BETA")
        for item in (
            {"content": json.dumps({"description": text, "content": text})},
            {"content": json.dumps({"concepts": [], "entities": []})},
        )
    ]
    with StreamTransport(responses) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        assert first.status == "added"
        before = storage_scope(root, identity)
        source.write_bytes(b"Broken workbook")
        failed = import_document(root, source, scope=before)
        assert failed.status == "failed"
        source.write_bytes(original)
        restored = import_document(root, source, scope=storage_scope(root, identity))
        assert restored.status == "added", restored.message
        assert {unit.name: unit.unit_id for unit in restored.units} == {
            unit.name: unit.unit_id for unit in first.units
        }
        assert len(storage_scope(root, identity).knowledge.unit_revisions()) == len(
            before.knowledge.unit_revisions()
        )
        assert len(service.requests) == 8


@pytest.mark.parametrize(
    "confirmation", ["normal", "receipt", "cleanup", "partial_cleanup", "cancelled"]
)
def test_workbook_publication_survives_next_unit_binding_failure_and_retry_does_not_resend(
    tmp_path, monkeypatch, confirmation
):
    from openpyxl import Workbook

    from ultrakb.storage.execution_store import ExecutionStore
    from ultrakb.storage.knowledge_store import KnowledgeStore
    from ultrakb.storage.usage_store import UsageStore

    root, identity, _ = question_knowledge(tmp_path)
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA"
    book.create_sheet("Beta")["A1"] = "BETA"
    source = tmp_path / "book.xlsx"
    book.save(source)
    published = False
    publish, bind = KnowledgeStore.publish, KnowledgeStore.bind
    finish = ExecutionStore.finish

    def confirm_then_fail(*args, **options):
        finish(*args, **options)
        raise OSError("Cleanup confirmation unavailable")

    def commit(*args, **options):
        nonlocal published
        result = publish(*args, **options)
        published = True
        return result

    def lose_binding(*args, **options):
        if published:
            raise OSError("Cannot bind the next worksheet yet")
        return bind(*args, **options)

    monkeypatch.setattr(KnowledgeStore, "publish", commit)
    monkeypatch.setattr(KnowledgeStore, "bind", lose_binding)
    responses = [
        item
        for text in ("ALPHA", "BETA")
        for item in (
            {"content": json.dumps({"description": text, "content": text})},
            {"content": json.dumps({"concepts": [], "entities": []})},
        )
    ]
    with StreamTransport(responses) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        if confirmation == "partial_cleanup":
            monkeypatch.setattr(ExecutionStore, "finish", confirm_then_fail)
        result = import_document(root, source, scope=storage_scope(root, identity))
        assert result.status == "partial", result.message
        alpha, beta = result.units
        assert alpha.status == "completed" and beta.status == "failed"
        assert alpha.knowledge_revision_id
        monkeypatch.setattr(KnowledgeStore, "bind", bind)
        monkeypatch.setattr(ExecutionStore, "finish", finish)
        if confirmation == "cleanup":
            monkeypatch.setattr(ExecutionStore, "finish", confirm_then_fail)
        if confirmation == "receipt":
            receipt = UsageStore.receipt

            def unavailable_receipt(*args, **options):
                if options.get("source_id") is not None:
                    raise OSError("Usage confirmation unavailable")
                return receipt(*args, **options)

            monkeypatch.setattr(UsageStore, "receipt", unavailable_receipt)
        source.unlink()
        cancelled = False

        def progress(event):
            nonlocal cancelled
            if (
                confirmation == "cancelled"
                and event.get("event") == "activity"
                and event["data"].get("stage") == "compile.generate.v1"
                and event["data"].get("status") == "started"
            ):
                cancelled = True

        context = ExecutionContext(cancelled=lambda: cancelled, on_event=progress)
        resumed = retry_source(
            root, result.source_id, context=context, scope=storage_scope(root, identity)
        )
        if confirmation == "cancelled":
            assert resumed.status == "stopped"
            assert (
                ExecutionStore(root, identity).status(context.storage_execution.execution_id)
                == "cancelled"
            )
            assert len(service.requests) == 2
            assert resumed.units[0].successful_revision_id == alpha.target_revision_id
            return
        assert resumed.status == "added", resumed.message
        assert resumed.units[0].target_revision_id == alpha.target_revision_id
        assert len(service.requests) == 4
        if confirmation != "receipt":
            assert resumed.model_usage["current"]["actual_sends"] == 2
            assert resumed.model_usage["cumulative"]["actual_sends"] == 4
        else:
            assert "confirm usage" in resumed.unfinished
        if confirmation == "cleanup":
            assert "confirm execution cleanup" in resumed.unfinished
        assert UsageStore(root, identity).summary(source_id=result.source_id)["actual_sends"] == 4
