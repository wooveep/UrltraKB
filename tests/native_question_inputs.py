"""Tiny real Office packages; no converted sources or model responses are supplied."""

from zipfile import ZIP_DEFLATED, ZipFile


def write_office_input(path):
    if path.suffix == ".docx":
        parts = {
            "[Content_Types].xml": (
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" '
                'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/word/document.xml" ContentType="application/vnd.'
                'openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
            ),
            "_rels/.rels": (
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                'relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.'
                'org/officeDocument/2006/relationships/officeDocument" '
                'Target="word/document.xml"/></Relationships>'
            ),
            "word/document.xml": (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/'
                'main"><w:body><w:p><w:r><w:t>The indoor room permits quiet reading.</w:t>'
                "</w:r></w:p><w:p><w:r><w:t>The garden room is closed.</w:t></w:r></w:p>"
                '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/></w:sectPr></w:body></w:document>'
            ),
        }
        with ZipFile(path, "w", ZIP_DEFLATED) as package:
            for name, content in parts.items():
                package.writestr(name, content)
    elif path.suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Briefing"
        slide.placeholders[1].text = "The visible-body code is BODY-17."
        slide.notes_slide.notes_text_frame.text = "The speaker-note code is NOTE-42."
        presentation.save(path)
    elif path.suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        public = workbook.active
        public.title = "Public"
        public["B2"] = "Quiet reading is forbidden."
        private = workbook.create_sheet("Private Budget")
        private["B2"] = "Quiet reading is permitted."
        private.sheet_state = "hidden"
        workbook.save(path)
        workbook.close()
    else:
        raise ValueError("Expected a DOCX, PPTX or XLSX test input")
    return path
