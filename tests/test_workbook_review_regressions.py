"""Source operations account for every sheet in the retained workbook inventory."""

import asyncio
from zipfile import ZipFile

import pytest

pytest_plugins = (
    "test_office_import",
    "block_fixtures",
    "workbook_fixtures",
    "processing_fixtures",
)


def test_embedded_workbook_recovery_requires_retry_for_uncreated_sheets(
    kb_dir, writer_document, three_sheets, pdf_model, monkeypatch, processing_only
):
    import ultrakb.application.ingestion as ingestion
    from ultrakb.application.pending import pending_status, process_pending, retry_pending_job
    from ultrakb.unit_publication import list_source_units, read_unit_publication

    with ZipFile(writer_document, "a") as package:
        package.writestr("word/embeddings/book.xlsx", three_sheets.read_bytes())
    processing_only(writer_document)
    process_pending(kb_dir, max_jobs=1)
    original = ingestion.process_import_unit
    count = 0

    def die_before_next_unit(*args, **kwargs):
        nonlocal count
        count += 1
        if count == 2:
            raise SystemExit("Worker termination between sheet publications")
        return original(*args, **kwargs)

    monkeypatch.setattr(ingestion, "process_import_unit", die_before_next_unit)
    with pytest.raises(SystemExit):
        process_pending(kb_dir, max_jobs=1)
    monkeypatch.setattr(ingestion, "process_import_unit", original)
    process_pending(kb_dir)
    job = next(row for row in pending_status(kb_dir)["jobs"] if row["kind"] == "import")
    units = list_source_units(kb_dir, job["source_id"])
    assert len(units) == 1
    view_id = "legacy"
    published = read_unit_publication(kb_dir, units[0].unit_id, view_id)
    assert job["status"] == "interrupted"
    assert process_pending(kb_dir)["processed"] == 0
    retry_pending_job(kb_dir, job["id"])
    process_pending(kb_dir)
    recovered = next(row for row in pending_status(kb_dir)["jobs"] if row["id"] == job["id"])
    assert recovered["status"] == "completed"
    assert len(list_source_units(kb_dir, job["source_id"])) == 3
    assert read_unit_publication(kb_dir, units[0].unit_id, view_id) == published


@pytest.mark.parametrize("delete_sheets", [False, True])
def test_all_confirmed_empty_sheets_refresh_retires_current_but_retains_history(
    kb_dir, three_sheets, delete_sheets
):
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook

    from ultrakb.application.refresh import refresh_knowledge_view, refresh_status
    from ultrakb.application.views import view_scope
    from ultrakb.unit_publication import read_head, wiki_versions

    first = publish_workbook(
        kb_dir,
        three_sheets,
    )
    assert first.status == "added"
    scope = view_scope(kb_dir, first.units[0].view_id)
    head = read_head(kb_dir, scope.view_id)
    history = view_scope(kb_dir, scope.view_id, historical_revision=head.knowledge_revision_id)
    old_pages = wiki_versions(kb_dir, history.wiki_dir)
    book = load_workbook(three_sheets)
    for sheet in book:
        sheet["A1"] = None
    if delete_sheets:
        del book["Beta"]
        del book["Gamma"]
    book.save(three_sheets)
    updated = publish_workbook(
        kb_dir,
        three_sheets,
    )
    assert updated.status == "added", updated.message
    assert len(updated.units) == 3
    assert all(unit.status in {"empty", "retired"} for unit in updated.units)
    status = refresh_status(kb_dir, scope=scope)
    assert not status["effective_inputs"] and status["needs_refresh"]
    result = asyncio.run(refresh_knowledge_view(kb_dir, scope=scope))
    assert result.status == "completed", result.message
    assert read_head(kb_dir, scope.view_id).inputs == {}
    assert not refresh_status(kb_dir, scope=scope)["needs_refresh"]
    assert not list((scope.wiki_dir / "summaries").glob("*.md"))
    assert wiki_versions(kb_dir, history.wiki_dir) == old_pages


def test_refresh_does_not_retire_unfinished_workbook_sheets(kb_dir, three_sheets):
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook

    from ultrakb.application.refresh import refresh_knowledge_view
    from ultrakb.application.views import view_scope
    from ultrakb.unit_publication import read_head, wiki_versions

    first = publish_workbook(
        kb_dir,
        three_sheets,
    )
    book = load_workbook(three_sheets)
    book["Alpha"]["A1"] = "ALPHA_UPDATED"
    book["Beta"]["A1"] = None
    book.save(three_sheets)
    with ZipFile(three_sheets) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet3.xml"] = b"<broken"
    with ZipFile(three_sheets, "w") as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    updated = publish_workbook(
        kb_dir,
        three_sheets,
    )
    assert updated.status == "partial"
    assert {unit.name: unit.status for unit in updated.units} == {
        "Alpha": "completed",
        "Beta": "empty",
        "Gamma": "failed",
    }
    scope = view_scope(kb_dir, first.units[0].view_id)
    head, pages = read_head(kb_dir, scope.view_id), wiki_versions(kb_dir, scope.wiki_dir)
    result = asyncio.run(refresh_knowledge_view(kb_dir, scope=scope))
    assert result.status == "blocked", result.message
    assert read_head(kb_dir, scope.view_id) == head
    assert wiki_versions(kb_dir, scope.wiki_dir) == pages
