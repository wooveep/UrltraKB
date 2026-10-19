"""Worksheet exclusions use complete retained inventory and exact publication receipts."""

import json

import pytest
from openpyxl import Workbook
from stream_transport import StreamTransport
from test_storage_questions import question_knowledge

from ultrakb.application.documents import import_document
from ultrakb.application.pages import read_page
from ultrakb.application.source_retry import retry_source
from ultrakb.application.storage_inventory import source_inventory
from ultrakb.application.views import storage_scope
from ultrakb.storage.catalog_store import CatalogStore
from ultrakb.storage.knowledge_validity import publication_validity
from ultrakb.storage.refresh_store import refresh_targets

pytestmark = pytest.mark.integration


def responses(*values):
    return [
        item
        for value in values
        for item in (
            {"content": json.dumps({"description": value, "content": value})},
            {"content": json.dumps({"concepts": [], "entities": []})},
        )
    ]


@pytest.mark.parametrize("retirement", ["empty", "withdrawn"])
def test_complete_inventory_retires_only_the_affected_sheet_and_retry_keeps_history(
    tmp_path, retirement
):
    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA"
    book.create_sheet("Beta")["A1"] = "BETA_HISTORICAL"
    book.save(source)
    with StreamTransport(responses("ALPHA", "BETA_HISTORICAL", "ALPHA")) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        assert first.status == "added", first.message
        before = storage_scope(root, identity)
        old = {unit.name: unit for unit in first.units}
        if retirement == "empty":
            book["Beta"]["A1"] = None
        else:
            del book["Beta"]
        book.save(source)
        second = import_document(root, source, scope=before)
        assert second.status == "added", second.message
        units = {unit.name: unit for unit in second.units}
        beta = units["Beta"]
        assert beta.status == ("empty" if retirement == "empty" else "retired")
        assert beta.unit_id == old["Beta"].unit_id
        assert beta.target_revision_id != old["Beta"].target_revision_id
        assert beta.model_usage["current"]["actual_sends"] == 0
        after = storage_scope(root, identity)
        assert second.publication_id == after.publication_id
        listed = next(row for row in source_inventory(after) if row["source_id"] == first.source_id)
        shown = next(unit for unit in listed["units"] if unit["unit_id"] == beta.unit_id)
        assert shown["status"] == beta.status
        assert shown["successful_revision_id"] == old["Beta"].target_revision_id
        targets = refresh_targets(after.knowledge)
        assert not targets.blocked
        assert {unit.unit_id for unit in targets.units} == {units["Alpha"].unit_id}
        reasons = publication_validity(after.knowledge)["refresh_reasons"]
        assert any(r["unit_id"] == beta.unit_id and r["kind"] == retirement for r in reasons)
        assert (
            "BETA_HISTORICAL"
            in read_page(root, f"summaries/{old['Beta'].doc_name}", scope=before).body
        )
        assert CatalogStore(root, identity).read_original(first.source_revision_id)
        source.unlink()
        repeated = retry_source(root, first.source_id, scope=after)
        assert repeated.status == "skipped", repeated.message
        assert {unit.unit_id: unit.target_revision_id for unit in repeated.units} == {
            unit.unit_id: unit.target_revision_id for unit in second.units
        }
        assert len(service.requests) == 6
    assert not (root / ".ultrakb/catalog").exists()
    assert not (root / ".ultrakb/hashes.json").exists()


def test_retirement_rechecks_the_source_target_inside_the_publication_commit(tmp_path, monkeypatch):
    from ultrakb.storage.coordination import storage_stage
    from ultrakb.storage.knowledge_store import KnowledgeStore

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA"
    book.create_sheet("Beta")["A1"] = "BETA"
    book.save(source)
    with StreamTransport(responses("ALPHA", "BETA", "ALPHA")) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        old_beta = next(unit for unit in first.units if unit.name == "Beta")
        del book["Beta"]
        book.save(source)
        publish, replacement, head = KnowledgeStore.publish, tmp_path / "replacement.xlsx", []
        book["Alpha"]["A1"] = "NEWER_SOURCE"
        book.save(replacement)

        def supersede_retirement(store, prepared, **options):
            with storage_stage(root, identity, read_only=True) as db:
                row = db.connection.execute(
                    "SELECT 1 FROM ukb_revision_source_changes WHERE revision_id = ? "
                    "AND unit_revision_id IS NOT NULL",
                    (prepared.revision_id,),
                ).fetchone()
            if row:
                head.append(store.bind().publication_id)
                CatalogStore(root, identity).admit(replacement, source_key=str(source.absolute()))
            return publish(store, prepared, **options)

        monkeypatch.setattr(KnowledgeStore, "publish", supersede_retirement)
        second = import_document(root, source, scope=storage_scope(root, identity))
        assert second.status == "partial", second.message
        beta = next(unit for unit in second.units if unit.unit_id == old_beta.unit_id)
        assert beta.status == "failed"
        assert CatalogStore(root, identity).unit_exclusion(beta.target_revision_id) is None
        assert storage_scope(root, identity).publication_id == head[0]
        assert len(service.requests) == 6


@pytest.mark.parametrize(
    "boundary", ["before_publish", "transaction", "lost_ack", "projection", "observation"]
)
def test_worksheet_retirement_recovers_without_resending_other_sheets(
    tmp_path, monkeypatch, boundary
):
    from ultrakb.storage.coordination import storage_stage
    from ultrakb.storage.execution_store import ExecutionStore
    from ultrakb.storage.knowledge_store import KnowledgeStore
    from ultrakb.storage.wiki_projection import ProjectionStore

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA"
    book.create_sheet("Beta")["A1"] = "BETA"
    book.save(source)
    with StreamTransport(responses("ALPHA", "BETA", "ALPHA")) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        old_beta = next(unit for unit in first.units if unit.name == "Beta")
        del book["Beta"]
        book.save(source)
        if boundary == "transaction":
            with storage_stage(root, identity) as db:
                db.connection.execute(
                    "CREATE TRIGGER interrupt_retirement BEFORE INSERT ON ukb_publication_receipts "
                    "WHEN EXISTS (SELECT 1 FROM ukb_revision_source_changes c "
                    "WHERE c.revision_id = NEW.revision_id AND c.unit_revision_id IS NOT NULL) "
                    "BEGIN SELECT RAISE(ABORT, 'retirement interrupted'); END"
                )
        publish, sync, finish = (
            KnowledgeStore.publish,
            ProjectionStore.sync,
            ExecutionStore.finish_stage,
        )
        committed = False

        def publish_boundary(store, prepared, **options):
            nonlocal committed
            with storage_stage(root, identity, read_only=True) as db:
                retirement = db.connection.execute(
                    "SELECT 1 FROM ukb_revision_source_changes WHERE revision_id = ? "
                    "AND unit_revision_id IS NOT NULL",
                    (prepared.revision_id,),
                ).fetchone()
            if retirement and boundary == "before_publish":
                raise OSError("Publication unavailable")
            result = publish(store, prepared, **options)
            if retirement:
                committed = True
                if boundary == "lost_ack":
                    raise OSError("Publication acknowledgement lost")
            return result

        def sync_boundary(store, **options):
            if committed and boundary == "projection":
                raise OSError("Projection unavailable")
            return sync(store, **options)

        def finish_boundary(store, *args, **options):
            if committed and boundary == "observation":
                raise OSError("Observation unavailable")
            return finish(store, *args, **options)

        with monkeypatch.context() as faults:
            faults.setattr(KnowledgeStore, "publish", publish_boundary)
            faults.setattr(ProjectionStore, "sync", sync_boundary)
            faults.setattr(ExecutionStore, "finish_stage", finish_boundary)
            second = import_document(root, source, scope=storage_scope(root, identity))
        interrupted = boundary in {"before_publish", "transaction"}
        assert second.status == ("partial" if interrupted else "added"), second.message
        beta = next(unit for unit in second.units if unit.unit_id == old_beta.unit_id)
        assert beta.status == ("failed" if interrupted else "retired")
        if boundary == "transaction":
            assert CatalogStore(root, identity).unit_exclusion(beta.target_revision_id) is None
            with storage_stage(root, identity) as db:
                db.connection.execute("DROP TRIGGER interrupt_retirement")
        if boundary == "projection":
            assert second.projection_pending
        assert len(service.requests) == 6
        source.unlink()
        result = retry_source(
            root,
            first.source_id,
            unit_id=beta.unit_id,
            scope=storage_scope(root, identity),
            source_revision_id=second.source_revision_id,
            unit_revision_id=beta.target_revision_id,
        )
        assert result.status == ("added" if interrupted else "skipped"), result.message
        assert result.units[0].status == "retired"
        assert not result.projection_pending
        assert not refresh_targets(storage_scope(root, identity).knowledge).blocked
        assert len(service.requests) == 6


def test_all_empty_sheets_refresh_without_model_work_and_keep_old_publication(tmp_path):
    import asyncio

    from ultrakb.application.refresh import refresh_knowledge_view

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "Alpha"
    book.active["A1"] = "ALPHA"
    book.save(source)
    with StreamTransport(responses("ALPHA")) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        before = storage_scope(root, identity)
        book["Alpha"]["A1"] = None
        book.save(source)
        empty = import_document(root, source, scope=before)
        assert empty.status == "added", empty.message
        after = storage_scope(root, identity)
        assert refresh_targets(after.knowledge).units == ()
        (root / ".env").unlink()
        (root / ".ultrakb/config.yaml").unlink()
        refreshed = asyncio.run(refresh_knowledge_view(root, scope=after))
        assert refreshed.status == "completed", refreshed.message
        final = storage_scope(root, identity)
        assert final.knowledge.unit_revisions() == ()
        assert final.knowledge.file_page(prefix="summaries/") == ()
        assert publication_validity(final.knowledge)["validity"] == "current"
        assert "ALPHA" in read_page(root, f"summaries/{first.units[0].doc_name}", scope=before).body
        assert len(service.requests) == 2
