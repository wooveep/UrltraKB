"""Native tables are compilation objects; cell positions remain source evidence."""

import json

import pytest

from openkb.application.documents import import_document
from tests.document_fixtures import write_docx
from tests.http_model_fixture import evidence_response


def write_table(path, rows):
    if path.suffix == ".docx":
        write_docx(
            path,
            "<w:tbl>"
            + "".join(
                "<w:tr>"
                + ("<w:trPr><w:tblHeader/></w:trPr>" if index == 0 else "")
                + "".join(f"<w:tc><w:p><w:r><w:t>{v}</w:t></w:r></w:p></w:tc>" for v in row)
                + "</w:tr>"
                for index, row in enumerate(rows)
            )
            + "</w:tbl>",
        )
    elif path.suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        table = slide.shapes.add_table(
            len(rows), len(rows[0]), Inches(1), Inches(1), Inches(7), Inches(4)
        ).table
        table.first_row = True
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        deck.save(path)
    elif path.suffix == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table

        book = Workbook()
        sheet = book.active
        sheet.title = "Limits"
        for row in rows:
            sheet.append(row)
        sheet.add_table(
            Table(displayName="Limits", ref=f"A1:{get_column_letter(len(rows[0]))}{len(rows)}")
        )
        book.save(path)
    else:
        import pymupdf

        doc = pymupdf.open()
        page = doc.new_page()
        for r in range(len(rows) + 1):
            page.draw_line((50, 50 + 30 * r), (450, 50 + 30 * r))
        for c in range(len(rows[0]) + 1):
            page.draw_line((50 + 200 * c, 50), (50 + 200 * c, 50 + 30 * len(rows)))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                page.insert_text((55 + 200 * c, 70 + 30 * r), value)
        doc.save(path)
        doc.close()


@pytest.mark.parametrize("suffix", [".docx", ".pdf", ".pptx", ".xlsx"])
def test_small_table_is_planned_in_one_document_request_without_fact_extraction(
    kb_dir, tmp_path, model_service, suffix
):
    path = tmp_path / ("limits" + suffix)
    rows = [("Parameter", "Required value"), ("timeout", "37 seconds"), ("retries", "3")]
    write_table(path, rows)
    requests = []

    def respond(body):
        payload = json.loads(body["messages"][-1]["content"])
        requests.append(payload)
        return evidence_response(payload)

    model_service.respond = respond
    result = import_document(kb_dir, path)
    assert result.knowledge_compilation == "completed", result
    assert not [request for request in requests if request["stage"] == "facts"]
    plans = [request for request in requests if request["stage"] == "planning"]
    assert len(plans) == 1
    assert {item["text"] for item in plans[0]["evidence"]["blocks"] if item["kind"] == "table"} == {
        value for row in rows for value in row
    }
    generated = [request for request in requests if request["stage"] == "generation"]
    assert len(generated) == 1
    assert {
        item["text"] for item in generated[0]["evidence"]["blocks"] if item["kind"] == "table"
    } == {value for row in rows for value in row}
    assert "table_objects" not in generated[0]


def test_large_spreadsheet_generation_keeps_complete_table_rows(kb_dir, tmp_path, model_service):
    from collections import Counter

    from openpyxl.utils.cell import coordinate_to_tuple

    path = tmp_path / "large.xlsx"
    rows = [("Name", "Limit", "Unit")] + [(f"item-{i}", str(i), "seconds") for i in range(45)]
    write_table(path, rows)
    generated = []

    def respond(body):
        payload = json.loads(body["messages"][-1]["content"])
        if payload["stage"] == "generation":
            generated.append(payload)
        return evidence_response(payload)

    model_service.respond = respond
    result = import_document(kb_dir, path)
    assert result.knowledge_compilation == "completed", result
    assert not result.omissions
    assert generated
    addresses = []
    for batch in generated:
        table_cells = [item for item in batch["evidence"]["blocks"] if item["kind"] == "table"]
        cells = [item["location"]["cell_address"] for item in table_cells]
        positions = [coordinate_to_tuple(cell) for cell in cells]
        assert set(Counter(row for row, col in positions).values()) == {3}
        addresses.extend(cells)
    assert len(addresses) == len(set(addresses)) == len(rows) * 3


def test_capacity_and_response_retry_do_not_split_a_row_or_merged_cell():
    from openkb.agent.evidence_retry import split_generation, split_generation_response

    def pair(row, column, rowspan=1):
        return (
            {"table_object": {"id": "table", "row": row, "column": column, "rowspan": rowspan}},
            {},
        )

    merged = [pair(1, 1, 2), pair(1, 2), pair(2, 2)]
    for split in (split_generation, split_generation_response):
        assert split(merged) == []
        assert split(merged + [pair(3, 1), pair(3, 2)]) == [merged, [pair(3, 1), pair(3, 2)]]


def test_unmarked_worksheet_regions_and_embedded_tables_keep_distinct_identity():
    from types import SimpleNamespace

    from openkb.agent.table_objects import source_table_objects
    from openkb.evidence import BlockDraft

    blocks = []
    for number, address in enumerate(["B3", "C3", "B4", "C4", "B9", "C9", "B10", "C10"]):
        draft = BlockDraft(
            "value",
            "table",
            {"kind": "xlsx", "sheet": "Sheet", "sheet_index": 1, "cell_address": address},
        )
        blocks.append(SimpleNamespace(id=str(number), **draft.__dict__))
    objects = source_table_objects(SimpleNamespace(blocks=blocks), SimpleNamespace(id="source"))
    assert {obj["position"]["range"] for obj in objects.values()} == {"B3:C4", "B9:C10"}
    assert len({obj["id"] for obj in objects.values()}) == 2

    for number in (1, 2):
        draft = BlockDraft(
            "same",
            "table",
            {
                "kind": "docx",
                "paragraph": 1,
                "attachment": {
                    "part": f"ole{number}.bin",
                    "name": "same.docx",
                    "blob": "a" * 64,
                    "position": {"kind": "docx", "table": 1, "row": 1, "cell": 1},
                },
            },
        )
        blocks.append(SimpleNamespace(id=f"embedded{number}", **draft.__dict__))
    objects = source_table_objects(SimpleNamespace(blocks=blocks), SimpleNamespace(id="source"))
    assert objects["embedded1"]["id"] != objects["embedded2"]["id"]


def test_large_table_stays_pending_when_a_whole_review_cannot_fit(kb_dir, tmp_path, model_service):
    import yaml

    config_path = kb_dir / ".openkb/config.yaml"
    config = yaml.safe_load(config_path.read_text())
    config["processing"].update(context_tokens=8192, output_tokens=512, max_requests=40)
    config_path.write_text(yaml.safe_dump(config))
    path = tmp_path / "scoped.docx"
    condition = "Only applicable to Linux version 7."
    rows = [("Name", "Value")] + [(f"limit-{i}", str(i)) for i in range(40)]
    write_docx(
        path,
        f"<w:p><w:r><w:t>{condition}</w:t></w:r></w:p><w:tbl>"
        + "".join(
            "<w:tr>"
            + "".join(f"<w:tc><w:p><w:r><w:t>{v}</w:t></w:r></w:p></w:tc>" for v in row)
            + "</w:tr>"
            for row in rows
        )
        + "</w:tbl>",
    )
    generated, reviewed = [], []

    def respond(body):
        payload = json.loads(body["messages"][-1]["content"])
        if payload["stage"] == "generation":
            generated.append(payload)
        if payload["stage"] == "verification":
            reviewed.append(payload)
        return evidence_response(payload)

    model_service.respond = respond
    result = import_document(kb_dir, path)
    assert result.knowledge_compilation == "completed", result
    # At this capacity the complete assembled page cannot receive its required
    # critical review. Preserve the generated source rows privately and report
    # a page-local continuation item; never substitute fragment verification.
    assert any(
        row["stage"] == "generation"
        and row["reason"] == "input_budget_exceeded"
        and row["items"] == ["concepts/notes"]
        for row in result.omissions
    ), result.omissions
    assert not (kb_dir / "wiki/concepts/notes.md").exists()
    assert generated
    assert not reviewed
    assert any(
        item["text"] == condition for batch in generated for item in batch["evidence"]["blocks"]
    )
    for batch in generated:
        rows_in_batch = {}
        for item in batch["evidence"]["blocks"]:
            if item["kind"] != "table":
                continue
            location = item["location"]
            rows_in_batch[location["row"]] = rows_in_batch.get(location["row"], 0) + 1
        assert set(rows_in_batch.values()) <= {2}


def test_document_table_uses_configured_context_headroom_before_row_splitting(
    kb_dir, tmp_path, model_service
):
    import litellm
    import yaml

    config_path = kb_dir / ".openkb/config.yaml"
    config = yaml.safe_load(config_path.read_text())
    config["processing"].update(context_tokens=4096, max_context_tokens=32768)
    config_path.write_text(yaml.safe_dump(config))
    path = tmp_path / "whole.docx"
    write_table(path, [("Parameter", "Value")] + [(f"limit-{i}", str(i)) for i in range(30)])
    result = import_document(kb_dir, path)
    assert result.knowledge_compilation == "completed", result
    assert not result.omissions
    generated = [
        call
        for call in model_service
        if json.loads(call["messages"][-1]["content"])["stage"] == "generation"
    ]
    assert len(generated) == 1
    request = generated[0]
    measured = (
        litellm.token_counter(model=request["model"], messages=request["messages"])
        + request["max_tokens"]
    )
    assert 4096 < measured <= 32768
