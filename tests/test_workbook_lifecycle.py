"""Workbook updates reconcile physical sheets and preserve per-unit publication facts."""

import pytest

pytest_plugins = ("test_office_import", "block_fixtures", "workbook_fixtures")


def test_xls_rename_and_reorder_use_the_same_identity_reconciliation(kb_dir, tmp_path):
    from pathlib import Path
    from shutil import copyfile

    from office_publication_fixtures import publish_workbook

    from ultrakb.documents import read_document_source

    fixtures = Path(__file__).parent / "fixtures/office"
    workbook = tmp_path / "workbook.xls"
    copyfile(fixtures / "typed-sheets.xls", workbook)
    first = publish_workbook(
        kb_dir,
        workbook,
    )
    original = {unit.name: unit.unit_id for unit in first.units}
    copyfile(fixtures / "renamed-sheets.xls", workbook)
    second = publish_workbook(
        kb_dir,
        workbook,
    )
    assert {unit.name: unit.unit_id for unit in second.units} == {
        "Renamed Alpha": original["Alpha"],
        "Beta": original["Beta"],
        "Gamma": original["Gamma"],
    }
    renamed = read_document_source(kb_dir, first.source_id, unit_id=original["Alpha"])
    assert renamed["sheet"]["ordinal"] == 2
    assert renamed["sheet"]["identity_basis"] == "unique_rename_in_complete_inventory"


def test_rename_and_reorder_keep_sheet_identity_without_using_table_position(
    kb_dir,
    three_sheets,
):
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook

    from ultrakb.documents import read_document_source

    first = publish_workbook(
        kb_dir,
        three_sheets,
    )
    original = {unit.name: unit.unit_id for unit in first.units}
    book = load_workbook(three_sheets)
    book["Alpha"].title = "Renamed Alpha"
    book.move_sheet("Gamma", offset=-2)
    book.save(three_sheets)
    second = publish_workbook(
        kb_dir,
        three_sheets,
    )
    assert second.status == "added", second.message
    current = {unit.name: unit.unit_id for unit in second.units}
    assert current == {
        "Renamed Alpha": original["Alpha"],
        "Beta": original["Beta"],
        "Gamma": original["Gamma"],
    }
    renamed = read_document_source(kb_dir, first.source_id, unit_id=original["Alpha"])
    assert renamed["sheet"]["ordinal"] == 2
    assert renamed["sheet"]["identity_basis"] == "unique_rename_in_complete_inventory"


def test_confirmed_empty_and_deleted_sheets_retire_only_their_contributions(
    kb_dir,
    three_sheets,
):
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook

    from ultrakb.application.refresh import refresh_status
    from ultrakb.application.sources import source_inventory
    from ultrakb.application.views import view_scope
    from ultrakb.documents import read_document_source

    first = publish_workbook(
        kb_dir,
        three_sheets,
    )
    units = {unit.name: unit for unit in first.units}
    book = load_workbook(three_sheets)
    book["Alpha"]["A1"] = "ALPHA_NEW_REVISION"
    book["Beta"]["A1"] = None
    del book["Gamma"]
    book.save(three_sheets)
    second = publish_workbook(
        kb_dir,
        three_sheets,
    )
    assert second.status == "added", second.message
    assert {unit.name: unit.status for unit in second.units} == {
        "Alpha": "completed",
        "Beta": "empty",
        "Gamma": "retired",
    }
    inventory = source_inventory(kb_dir)[0]
    assert len(inventory["units"]) == 3
    empty = read_document_source(kb_dir, first.source_id, unit_id=units["Beta"].unit_id)
    assert empty["status"] == "completed" and "BETA_SHEET" in empty["content"]
    assert next(unit for unit in inventory["units"] if unit["name"] == "Beta")["status"] == "empty"
    retired = read_document_source(kb_dir, first.source_id, unit_id=units["Gamma"].unit_id)
    assert retired["validity"] == "needs_refresh" and retired["status"] == "completed"
    assert "GAMMA_SHEET" in retired["content"]
    assert (
        next(unit for unit in inventory["units"] if unit["name"] == "Gamma")["status"] == "retired"
    )
    assert retired["target_source_revision_id"] == second.source_revision_id
    pending = refresh_status(kb_dir, scope=view_scope(kb_dir, first.units[0].view_id))[
        "needs_refresh"
    ]
    assert any(
        reason["kind"] == "empty" for reason in pending[f"summaries/{units['Beta'].doc_name}.md"]
    )
    assert any(
        reason["kind"] == "withdrawn"
        for reason in pending[f"summaries/{units['Gamma'].doc_name}.md"]
    )


def test_failed_and_objects_only_sheets_do_not_become_successful_empty_targets(
    kb_dir, three_sheets, tmp_path
):
    from zipfile import ZipFile

    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image
    from PIL import Image as PillowImage

    from ultrakb.application.sources import source_inventory
    from ultrakb.documents import read_document_source

    first = publish_workbook(
        kb_dir,
        three_sheets,
    )
    original = {unit.name: unit for unit in first.units}
    book = load_workbook(three_sheets)
    book["Alpha"]["A1"] = "ALPHA_NEW"
    book["Beta"]["A1"] = None
    picture = tmp_path / "red.png"
    PillowImage.new("RGB", (10, 10), "red").save(picture)
    book["Beta"].add_image(Image(picture), "B2")
    book.save(three_sheets)
    with ZipFile(three_sheets) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet3.xml"] = b"<broken"
    with ZipFile(three_sheets, "w") as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    second = publish_workbook(
        kb_dir,
        three_sheets,
    )
    assert second.status == "partial", second.message
    inventory = source_inventory(kb_dir)[0]
    units = {unit["name"]: unit for unit in inventory["units"]}
    assert units["Beta"]["content_state"] == "objects_only"
    assert units["Gamma"]["content_state"] == "parse_failed"
    for name in ("Beta", "Gamma"):
        assert units[name]["status"] == {"Beta": "blocked", "Gamma": "failed"}[name]
        assert units[name]["successful_source_revision_id"] == first.source_revision_id
    old = read_document_source(kb_dir, first.source_id, unit_id=original["Gamma"].unit_id)
    assert old["source_revision_id"] == first.source_revision_id
    assert "GAMMA_SHEET" in old["content"]


def test_workbook_removal_preview_names_owned_sheets(kb_dir, three_sheets):
    from office_publication_fixtures import publish_workbook

    from ultrakb.application.removal import preview_removal

    first = publish_workbook(kb_dir, three_sheets)
    preview = preview_removal(kb_dir, first.source_id)
    labels = str(preview.plan.actions)
    assert all(unit.doc_name in labels for unit in first.units)


@pytest.mark.parametrize("keep_anchor", [False, True])
def test_all_renamed_and_reordered_cannot_inherit_another_sheets_history(
    kb_dir, three_sheets, keep_anchor
):
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook

    first = publish_workbook(
        kb_dir,
        three_sheets,
    )
    original = {unit.unit_id for unit in first.units}
    book = load_workbook(three_sheets)
    for sheet in book:
        if not keep_anchor or sheet.title != "Alpha":
            sheet.title = "New " + sheet.title
    book.move_sheet("New Gamma", offset=-1 if keep_anchor else -2)
    book.save(three_sheets)
    second = publish_workbook(
        kb_dir,
        three_sheets,
    )
    current = {unit.unit_id for unit in second.units if unit.name.startswith("New ")}
    assert len(current) == (2 if keep_anchor else 3) and not current & original


def test_deleted_sheet_is_withdrawn(kb_dir, three_sheets):
    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook

    from ultrakb.documents import read_document_source

    first = publish_workbook(
        kb_dir,
        three_sheets,
    )
    gamma = next(u for u in first.units if u.name == "Gamma")
    book = load_workbook(three_sheets)
    del book["Gamma"]
    book.save(three_sheets)
    second = publish_workbook(
        kb_dir,
        three_sheets,
    )
    read = read_document_source(kb_dir, first.source_id, unit_id=gamma.unit_id)
    assert read["target_source_revision_id"] == second.source_revision_id
    assert read["status"] == "completed" and "GAMMA_SHEET" in read["content"]
    from ultrakb.application.sources import source_inventory

    assert (
        next(
            unit
            for unit in source_inventory(kb_dir)[0]["units"]
            if unit["unit_id"] == gamma.unit_id
        )["status"]
        == "retired"
    )
    assert read["validity"] == "needs_refresh"


def test_unextracted_background_picture_is_not_confirmed_empty(kb_dir, three_sheets):
    from io import BytesIO
    from zipfile import ZipFile

    from office_publication_fixtures import publish_workbook
    from openpyxl import load_workbook
    from PIL import Image

    from ultrakb.application.sources import source_inventory

    publish_workbook(
        kb_dir,
        three_sheets,
    )
    book = load_workbook(three_sheets)
    book["Alpha"]["A1"] = None
    book.save(three_sheets)
    with ZipFile(three_sheets) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b"</worksheet>",
        b'<picture xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        b'r:id="rBg"/></worksheet>',
    )
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rBg" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        b'Target="/xl/media/background.png"/></Relationships>'
    )
    image = BytesIO()
    Image.new("RGB", (10, 10), "red").save(image, format="PNG")
    parts["xl/media/background.png"] = image.getvalue()
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>", b'<Default Extension="png" ContentType="image/png"/></Types>'
    )
    with ZipFile(three_sheets, "w") as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    second = publish_workbook(
        kb_dir,
        three_sheets,
    )
    alpha = next(unit for unit in second.units if unit.name == "Alpha")
    assert alpha.status != "empty", (alpha, source_inventory(kb_dir))
