"""Office evidence retains real original positions through publication and question tools."""

import asyncio

import pytest

from openkb.application.conversations import ask_question, continue_conversation
from openkb.application.documents import import_document
from tests.http_model_fixture import original_source_answer


def test_spreadsheet_equal_values_retain_cell_and_header_identity_in_query_and_chat(
    kb_dir, tmp_path, model_service
):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table

    book = Workbook()
    sheet = book.active
    sheet.title = "Timeouts"
    sheet.append(["Connection timeout", "Retry interval"])
    sheet.append(["30 seconds", "30 seconds"])
    sheet.add_table(Table(displayName="Limits", ref="A1:B2"))
    source = tmp_path / "settings.xlsx"
    book.save(source)
    imported = import_document(kb_dir, source)
    assert imported.knowledge_compilation == "completed", imported
    views = []

    def inspect(tree, result):
        values = [row for row in result["evidence"] if row["text"] == "30 seconds"]
        assert len(values) == 2
        assert {row["location"]["cell_address"] for row in values} == {"A2", "B2"}
        assert {row["location"]["sheet"] for row in values} == {"Timeouts"}
        assert "Connection timeout" in values[0]["context"]
        assert "Retry interval" in values[1]["context"]
        assert values[0]["reference"]["block_id"] != values[1]["reference"]["block_id"]
        assert all(row["reference"]["parse_id"] == imported.parse_id for row in values)
        views.append(values)
        return "Connection timeout: 30 seconds (A2). Retry interval: 30 seconds (B2)."

    model_service.chat_response = original_source_answer(inspect)
    for operation in (ask_question, continue_conversation):
        answer = asyncio.run(operation(kb_dir, "Explain both timeouts and cite the cells."))
        assert answer.status == "completed", answer
        assert "A2" in answer.answer and "B2" in answer.answer
    assert len(views) == 2


@pytest.mark.parametrize("asset_state", ["valid", "missing", "corrupt", "symlink_loop"])
def test_slides_repeated_text_retains_object_identity_and_original_image(
    kb_dir, tmp_path, model_service, asset_state
):
    import io

    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    slides = Presentation()
    image = io.BytesIO()
    Image.new("RGB", (24, 24), "blue").save(image, format="PNG")
    for title in ("Connection", "Retry"):
        slide = slides.slides.add_slide(slides.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        box.text = "30 seconds"
        image.seek(0)
        slide.shapes.add_picture(image, Inches(1), Inches(3), width=Inches(1))
    source = tmp_path / "timeouts.pptx"
    slides.save(source)
    imported = import_document(kb_dir, source)
    assert imported.knowledge_compilation == "completed", imported
    if asset_state != "valid":
        for published in (kb_dir / "wiki/sources/images").glob("*"):
            if asset_state in {"missing", "symlink_loop"}:
                published.unlink()
                if asset_state == "symlink_loop":
                    published.symlink_to(published.name)
            else:
                published.write_bytes(b"damaged image")
    views = []

    def inspect(tree, result):
        values = [row for row in result["evidence"] if row["text"] == "30 seconds"]
        assert len(values) == 2
        assert {row["location"]["slide"] for row in values} == {1, 2}
        assert all(row["location"]["object_id"] == 3 for row in values)
        assert all("bbox" in row["location"] for row in values)
        assert "Connection" in values[0]["context"] and "Retry" in values[1]["context"]
        pictures = [row for row in result["evidence"] if row["kind"] == "image"]
        assert len(pictures) == 2 and all(row["assets"] for row in pictures)
        for picture in pictures:
            if asset_state != "valid":
                assert picture["images"] == []
                continue
            assert len(picture["images"]) == 1
            figure = picture["images"][0]
            assert figure["asset"] in picture["assets"]
            assert (kb_dir / "wiki" / figure["path"]).read_bytes() == image.getvalue()
            assert figure["markdown"] == f"![原图]({figure['path']})"
        assert all(row["reference"]["parse_id"] == imported.parse_id for row in values)
        views.append(values)
        return "Connection: 30 seconds (slide 1). Retry: 30 seconds (slide 2).\n" + (
            pictures[0]["images"][0]["markdown"] if asset_state == "valid" else "Image unavailable."
        )

    model_service.chat_response = original_source_answer(inspect)
    for operation in (ask_question, continue_conversation):
        answer = asyncio.run(operation(kb_dir, "Compare the timeouts on both slides."))
        assert answer.status == "completed", answer
        assert "slide 1" in answer.answer and "slide 2" in answer.answer
        assert ("![原图](sources/images/" in answer.answer) == (asset_state == "valid")
        if operation is continue_conversation:
            from openkb.agent.chat_session import load_session

            saved = load_session(kb_dir, answer.session_id)
            assert saved.assistant_texts[-1] == answer.answer
    assert len(views) == 2


def test_docx_single_cell_prose_is_not_its_own_header(kb_dir, tmp_path, model_service):
    from tests.document_fixtures import write_docx

    source = tmp_path / "prose.docx"
    write_docx(
        source,
        "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>"
        "Timeout is 42 seconds.</w:t></w:r></w:p></w:tc></w:tr></w:tbl>",
    )
    imported = import_document(kb_dir, source)
    assert imported.knowledge_compilation == "completed", imported

    def inspect(tree, result):
        row = next(row for row in result["evidence"] if "42 seconds" in row["text"])
        assert row["location"]["kind"] == "docx"
        assert row["location"]["table"] == 1 and row["location"]["cell"] == 1
        assert "header: Timeout" not in row["context"]
        return "Timeout is 42 seconds."

    model_service.chat_response = original_source_answer(inspect)
    assert asyncio.run(ask_question(kb_dir, "What is the timeout?")).status == "completed"


def test_spreadsheet_merged_subject_reaches_each_original_row(kb_dir, tmp_path, model_service):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Ports"
    sheet.append(["Service", "Port", "Direction"])
    sheet.append(["nacos", 31848, "external"])
    sheet.append([None, 8848, "internal"])
    sheet.merge_cells("A2:A3")
    source = tmp_path / "ports.xlsx"
    book.save(source)
    imported = import_document(kb_dir, source)
    assert imported.knowledge_compilation == "completed", imported

    def inspect(tree, result):
        row = next(row for row in result["evidence"] if row["text"] == "8848")
        assert row["location"]["cell_address"] == "B3"
        assert "nacos" in row["context"] and "A2:A3" in row["context"]
        assert "internal" in row["context"]
        assert "header role unconfirmed" in row["context"]
        return "nacos internal port: 8848 (B3)."

    model_service.chat_response = original_source_answer(inspect)
    assert (
        asyncio.run(ask_question(kb_dir, "What is the internal nacos port?")).status == "completed"
    )


def test_original_read_preserves_stored_crlf_offsets(kb_dir, tmp_path, model_service):
    from openkb.application.source_actions import read_source_evidence
    from openkb.evidence import Evidence, ParseStore
    from tests.document_fixtures import write_docx

    path = tmp_path / "exact.docx"
    write_docx(path, "<w:p><w:r><w:t>alpha&#13;&#10;Timeout: 42 seconds.</w:t></w:r></w:p>")
    imported = import_document(kb_dir, path)
    assert imported.knowledge_compilation == "completed", imported
    parsed = ParseStore(kb_dir).load(imported.parse_id)
    block = parsed.blocks[0]
    reference = Evidence(imported.source_id, imported.input_version, parsed.id, block.id)
    text = read_source_evidence(kb_dir, reference, max_chars=4000).text
    assert text == "alpha\r\nTimeout: 42 seconds."
    tail = Evidence(imported.source_id, imported.input_version, parsed.id, block.id, 7)
    assert read_source_evidence(kb_dir, tail, max_chars=100).text == "Timeout: 42 seconds."


def test_docx_numeric_style_uses_declared_outline_not_its_name(kb_dir, tmp_path, model_service):
    from openkb.application.source_history import source_status
    from tests.document_fixtures import write_docx

    path = tmp_path / "outline.docx"
    write_docx(
        path,
        '<w:p><w:pPr><w:pStyle w:val="7"/></w:pPr><w:r><w:t>Retry procedure</w:t></w:r></w:p>'
        "<w:p><w:r><w:t>Wait 42 seconds.</w:t></w:r></w:p>",
        styles='<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:style w:type="paragraph" w:styleId="7"><w:name w:val="7"/>'
        '<w:pPr><w:outlineLvl w:val="3"/></w:pPr></w:style></w:styles>',
    )
    imported = import_document(kb_dir, path)
    assert imported.knowledge_compilation == "completed", imported
    nav = source_status(kb_dir, imported.source_id)["navigation"]
    assert any(
        node["title"] == "Retry procedure" and node["structure_origin"] == "native"
        for node in nav["nodes"]
    )


def test_long_spreadsheet_context_is_pageable_in_published_question_tools(
    kb_dir, tmp_path, model_service
):
    import json

    from agents.tool_context import ToolContext
    from openpyxl import Workbook

    from openkb.agent.source_tools import source_tools
    from openkb.application.source_history import source_status

    book = Workbook()
    sheet = book.active
    sheet.append(["Timeout is 42 seconds.", "x" * 17000])
    path = tmp_path / "long-context.xlsx"
    book.save(path)
    imported = import_document(kb_dir, path)
    assert imported.knowledge_compilation == "completed", imported
    nav = source_status(kb_dir, imported.source_id)["navigation"]
    tool = next(tool for tool in source_tools(kb_dir)[0] if tool.name == "read_source_node")
    args = {
        "source_id": imported.source_id,
        "node_id": nav["nodes"][0]["id"],
        "offset": 1,
        "start": 0,
        "max_chars": 16000,
    }
    read = asyncio.run(
        tool.on_invoke_tool(
            ToolContext(
                context=None, tool_name="read_source_node", tool_call_id="test", tool_arguments="{}"
            ),
            json.dumps(args),
        )
    )
    first = json.loads(read)
    assert first["evidence"][0]["text"] == "Timeout is 42 seconds."
    assert first["next"] and not first["evidence"][0]["context_complete"]
    args.update(first["next"])
    second = json.loads(
        asyncio.run(
            tool.on_invoke_tool(
                ToolContext(
                    context=None,
                    tool_name="read_source_node",
                    tool_call_id="test",
                    tool_arguments="{}",
                ),
                json.dumps(args),
            )
        )
    )
    assert second["evidence"][0]["context_only"]
    assert second["evidence"][0]["context_complete"]
