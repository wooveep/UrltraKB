"""Explicit processing requests use managed inputs and preserve their recovery identity."""

import json

import pytest
from stream_transport import StreamTransport
from test_storage_questions import question_knowledge

from ultrakb.application.documents import import_document
from ultrakb.application.execution import ExecutionContext
from ultrakb.application.knowledge_tasks import read_knowledge_execution, read_knowledge_task
from ultrakb.application.reprocessing import (
    ReprocessingConflict,
    preview_reprocessing,
    reprocess_source,
)
from ultrakb.application.views import storage_scope
from ultrakb.config import load_config, save_config
from ultrakb.storage.catalog_store import CatalogStore
from ultrakb.storage.knowledge_store import KnowledgeStore

pytestmark = pytest.mark.integration


def responses():
    return [
        {"content": json.dumps({"description": "Value", "content": "The value is 17."})},
        {"content": json.dumps({"concepts": [], "entities": []})},
    ]


def test_reviewed_reprocessing_creates_a_processing_revision_from_the_managed_original(tmp_path):
    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.txt"
    source.write_text("The value is 17.")
    with StreamTransport(responses() * 3) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        assert first.status == "added", first.message
        source.unlink()
        public = storage_scope(root, identity)
        preview = preview_reprocessing(root, first.source_id, scope=public)
        assert preview["status"] == "ready", preview
        assert preview["original"]["available"]
        assert preview["units"][0]["target_revision_id"] == first.units[0].target_revision_id
        assert len(service.requests) == 2
        second = reprocess_source(root, first.source_id, version=preview["version"], scope=public)
        assert second.status == "added", second.message
        assert second.source_revision_id == first.source_revision_id
        assert second.units[0].unit_id == first.units[0].unit_id
        assert second.units[0].target_revision_id != first.units[0].target_revision_id
        assert len(service.requests) == 4
        repeated = reprocess_source(
            root,
            first.source_id,
            version=preview["version"],
            scope=storage_scope(root, identity),
        )
        assert repeated.status == "skipped", repeated.message
        assert len(service.requests) == 4
        assert repeated.model_usage["current"]["actual_sends"] == 0
        assert repeated.model_usage["cumulative"]["actual_sends"] == 4
        current = storage_scope(root, identity)
        next_preview = preview_reprocessing(root, first.source_id, scope=current)
        assert not next_preview["resume"]
        assert next_preview["version"] != preview["version"]
        third = reprocess_source(
            root, first.source_id, version=next_preview["version"], scope=current
        )
        assert third.status == "added", third.message
        with pytest.raises(ReprocessingConflict, match="target|preview"):
            reprocess_source(
                root,
                first.source_id,
                version=preview["version"],
                scope=storage_scope(root, identity),
            )
        assert len(service.requests) == 6
    catalog = CatalogStore(root, identity)
    assert catalog.read_original(first.source_revision_id) == b"The value is 17."
    assert first.units[0].target_revision_id in public.knowledge.unit_revisions()
    assert not (root / ".ultrakb/catalog").exists()


def test_reprocessing_rejects_changed_reviewed_policy_without_changing_the_target(tmp_path):
    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.txt"
    source.write_text("The value is 17.")
    with StreamTransport(responses()) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        public = storage_scope(root, identity)
        preview = preview_reprocessing(root, first.source_id, scope=public)
        path = root / ".ultrakb/config.yaml"
        config = load_config(path)
        save_config(
            path,
            {**config, "model_capacity": {**config["model_capacity"], "max_input_tokens": 32768}},
        )
        with pytest.raises(ReprocessingConflict, match="preview|policy"):
            reprocess_source(root, first.source_id, version=preview["version"], scope=public)
        assert len(service.requests) == 2
        target = CatalogStore(root, identity).unit(first.units[0].unit_id)
        assert target.unit_revision_id == first.units[0].target_revision_id


@pytest.mark.parametrize(
    "resume_mode", ["same-kb", "restored-generation", "external-publication", "external-retry"]
)
def test_reprocessing_resumes_accepted_compilation_after_publication_failure(
    tmp_path, monkeypatch, resume_mode
):
    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.txt"
    source.write_text("The value is 17.")
    with StreamTransport(responses() * 2) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        public = storage_scope(root, identity)
        preview = preview_reprocessing(root, first.source_id, scope=public)
        with monkeypatch.context() as crash:

            def fail_publication(*args, **kwargs):
                raise OSError("interrupted before publication")

            crash.setattr(KnowledgeStore, "publish", fail_publication)
            failed_context = ExecutionContext()
            failed = reprocess_source(
                root,
                first.source_id,
                version=preview["version"],
                scope=public,
                context=failed_context,
            )
        assert failed.status == "blocked", failed.message
        assert len(service.requests) == 4
        target = CatalogStore(root, identity).unit(first.units[0].unit_id)
        if resume_mode in {"external-publication", "external-retry"}:
            from ultrakb.application.pages import read_page, save_page

            path = f"summaries/{first.units[0].doc_name}.md"
            page = read_page(root, path, scope=public)
            saved = save_page(
                root, path, page.body + "\nHuman annotation.\n", version=page.version, scope=public
            )
            assert saved.status == "saved"
            with pytest.raises(ReprocessingConflict, match="publication|preview|Knowledge"):
                if resume_mode == "external-retry":
                    from ultrakb.application.source_retry import retry_source

                    retry_source(root, first.source_id, scope=storage_scope(root, identity))
                else:
                    reprocess_source(
                        root,
                        first.source_id,
                        version=preview["version"],
                        scope=storage_scope(root, identity),
                    )
            assert len(service.requests) == 4
            return
        if resume_mode == "restored-generation":
            from ultrakb.application.storage_maintenance import backup_knowledge, restore_knowledge

            backup = tmp_path / "backup"
            backup_knowledge(root, backup, scope=public)
            root = tmp_path / "restored"
            restored = restore_knowledge(backup, root)
            assert restored.identity.kb_id == identity.kb_id
            assert restored.identity.generation != identity.generation
            identity = restored.identity
            public = storage_scope(root, identity)
        recovered = reprocess_source(
            root, first.source_id, version=preview["version"], scope=public
        )
        assert recovered.status == "added", recovered.message
        assert recovered.units[0].target_revision_id == target.unit_revision_id
        assert recovered.model_usage["current"]["actual_sends"] == 0
        assert len(service.requests) == 4


def test_workbook_reprocessing_reparses_into_an_independent_inventory(tmp_path):
    from openpyxl import Workbook

    from ultrakb.storage.workbook_store import WorkbookStore

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    workbook.active.title = "Value"
    workbook.active["A1"] = "The value is 17."
    workbook.create_sheet("Empty")
    workbook.save(source)
    with StreamTransport(responses() * 2) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        assert {unit.status for unit in first.units} == {"completed", "empty"}
        source.unlink()
        before = {unit.unit_id: unit.target_revision_id for unit in first.units}
        public = storage_scope(root, identity)
        preview = preview_reprocessing(root, first.source_id, scope=public)
        assert preview["range"] == "all_worksheets"
        assert preview["status"] == "ready", preview
        second = reprocess_source(root, first.source_id, version=preview["version"], scope=public)
        assert {unit.status for unit in second.units} == {"completed", "empty"}, second.message
        assert set(before) == {unit.unit_id for unit in second.units}
        assert all(unit.target_revision_id != before[unit.unit_id] for unit in second.units)
        inventories = WorkbookStore(root, identity)
        original = inventories.read(first.source_revision_id)
        new = inventories.read(first.source_revision_id, processing_key=preview["version"])
        assert original is not None and new is not None
        assert [(sheet.key, sheet.cells, sheet.content_state) for sheet in new.sheets] == [
            (sheet.key, sheet.cells, sheet.content_state) for sheet in original.sheets
        ]
        assert all(
            sheet.identity_from_revision_id == first.source_revision_id for sheet in new.sheets
        )
        assert len(service.requests) == 4


def test_reprocessing_an_unreadable_original_stays_retained_without_model_work(tmp_path):
    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.txt"
    source.write_text(" \n")
    public = storage_scope(root, identity)
    with StreamTransport(responses()) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=public)
        source.unlink()
        preview = preview_reprocessing(root, first.source_id, scope=public)
        assert preview["status"] == "ready", preview
        second = reprocess_source(root, first.source_id, version=preview["version"], scope=public)
        assert second.status == "blocked", second.message
        assert second.units[0].target_revision_id != first.units[0].target_revision_id
        target = CatalogStore(root, identity).unit(second.units[0].unit_id)
        assert (
            CatalogStore(root, identity).processing(target.processing_id).content_state == "empty"
        )
        assert storage_scope(root, identity).publication_id == public.publication_id
        assert not service.requests


def test_reprocessing_guards_a_preview_with_no_processing_unit(tmp_path, monkeypatch):
    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "blank.txt"
    source.write_text(" \n")
    catalog = CatalogStore(root, identity)
    admitted = catalog.admit(source, source_key=str(source.absolute()))
    public = storage_scope(root, identity)
    competing, retain = [], CatalogStore.retain_processing

    def concurrent_processing(self, *args, **kwargs):
        value = retain(self, *args, **kwargs)
        if not competing:
            competing.append(None)
            competing[0] = import_document(root, source, scope=public)
        return value

    with StreamTransport(responses()) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        preview = preview_reprocessing(root, admitted.source_id, scope=public)
        assert not preview["units"]
        monkeypatch.setattr(CatalogStore, "retain_processing", concurrent_processing)
        result = reprocess_source(
            root, admitted.source_id, version=preview["version"], scope=public
        )
        assert result.status == "blocked"
        target = catalog.unit(competing[0].units[0].unit_id)
        assert target.unit_revision_id == competing[0].units[0].target_revision_id
        assert target.generation == 1
        assert "changed" in result.message
        refreshed = preview_reprocessing(root, admitted.source_id, scope=public)
        assert not refreshed["resume"]
        assert refreshed["version"] != preview["version"]
        continued = reprocess_source(
            root, admitted.source_id, version=refreshed["version"], scope=public
        )
        assert continued.status == "blocked"
        assert continued.units[0].target_revision_id != target.unit_revision_id
        assert catalog.unit(target.unit_id).generation == 2
        assert not service.requests


def test_workbook_reprocessing_resumes_its_own_partial_publication_chain(tmp_path, monkeypatch):
    from openpyxl import Workbook

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "The value is 17."
    workbook.create_sheet("Other")["A1"] = "The value is 17."
    workbook.save(source)
    with StreamTransport(responses() * 4) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        assert first.status == "added", first.message
        source.unlink()
        public = storage_scope(root, identity)
        preview = preview_reprocessing(root, first.source_id, scope=public)
        publish, count = KnowledgeStore.publish, 0

        def fail_second(self, *args, **kwargs):
            nonlocal count
            count += 1
            if count == 2:
                raise OSError("second worksheet publication interrupted")
            return publish(self, *args, **kwargs)

        with monkeypatch.context() as crash:
            crash.setattr(KnowledgeStore, "publish", fail_second)
            partial = reprocess_source(
                root, first.source_id, version=preview["version"], scope=public
            )
        assert partial.status == "partial", partial.message
        assert len(service.requests) == 8
        resumed = reprocess_source(
            root, first.source_id, version=preview["version"], scope=storage_scope(root, identity)
        )
        assert resumed.status == "added", resumed.message
        assert all(unit.status == "completed" for unit in resumed.units)
        assert resumed.model_usage["current"]["actual_sends"] == 0
        assert len(service.requests) == 8


@pytest.mark.parametrize("via_retry", [False, True], ids=["reprocessing", "source-retry"])
def test_workbook_reprocessing_retries_a_failed_inventory_without_losing_its_budget(
    tmp_path, monkeypatch, via_retry
):
    from openpyxl import Workbook

    from ultrakb.workbooks import xlsx

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "The value is 17."
    workbook.save(source)
    with StreamTransport(responses() * 2) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        source.unlink()
        public = storage_scope(root, identity)
        preview = preview_reprocessing(root, first.source_id, scope=public)
        parse, attempts = xlsx.read_xlsx, []

        def fail_once(path):
            attempts.append(path)
            if len(attempts) == 1:
                raise OSError("temporary parser read failure")
            return parse(path)

        monkeypatch.setattr(xlsx, "read_xlsx", fail_once)
        context = ExecutionContext()
        failed = reprocess_source(
            root, first.source_id, version=preview["version"], scope=public, context=context
        )
        assert failed.status == "failed", failed.message
        assert len(service.requests) == 2
        detail = read_knowledge_execution(
            root, context.storage_execution.execution_id, scope=public
        )
        stage = next(
            stage for stage in detail.stages if stage.stage_code == "workbook.inventory.v1"
        )
        assert stage.state == "failed"
        assert stage.error_code == "WorkbookInventoryFailed"
        repeated_preview = preview_reprocessing(root, first.source_id, scope=public)
        assert repeated_preview["version"] == preview["version"]
        continued = ExecutionContext()
        if via_retry:
            from ultrakb.application.source_retry import retry_source

            recovered = retry_source(root, first.source_id, scope=public, context=continued)
        else:
            recovered = reprocess_source(
                root, first.source_id, version=preview["version"], scope=public, context=continued
            )
        assert recovered.status == "added", recovered.message
        assert len(attempts) == 2
        assert len(service.requests) == 4
        resumed = read_knowledge_execution(
            root, continued.storage_execution.execution_id, scope=public
        )
        assert resumed.task.task_id == detail.task.task_id
        assert resumed.task.budget.model_budget_id == detail.task.budget.model_budget_id
        assert resumed.task.budget.used == 2


def test_two_concurrent_resumes_cannot_start_two_executions_for_one_request(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier, Event, Lock

    from ultrakb.storage.execution_store import ExecutionStore

    root, identity, _ = question_knowledge(tmp_path)
    source = tmp_path / "input.txt"
    source.write_text("The value is 17.")
    with StreamTransport(responses() * 2) as service:
        (root / ".env").write_text(f"LLM_API_KEY=offline-test-key\nOPENAI_API_BASE={service.url}\n")
        first = import_document(root, source, scope=storage_scope(root, identity))
        public = storage_scope(root, identity)
        preview = preview_reprocessing(root, first.source_id, scope=public)
        with monkeypatch.context() as crash:

            def fail_publication(*args, **kwargs):
                raise OSError("publication interrupted")

            crash.setattr(KnowledgeStore, "publish", fail_publication)
            failed_context = ExecutionContext()
            failed = reprocess_source(
                root,
                first.source_id,
                version=preview["version"],
                scope=public,
                context=failed_context,
            )
        assert failed.status == "blocked"
        begin = ExecutionStore.begin
        rendezvous, created, second_done = Barrier(2), Event(), Event()
        guard, arrivals = Lock(), []

        def simultaneous_begin(self, *args, **kwargs):
            rendezvous.wait(timeout=10)
            with guard:
                ordinal = len(arrivals)
                arrivals.append(ordinal)
            if ordinal == 0:
                execution = begin(self, *args, **kwargs)
                created.set()
                assert second_done.wait(10)
                return execution
            assert created.wait(10)
            try:
                return begin(self, *args, **kwargs)
            finally:
                second_done.set()

        monkeypatch.setattr(ExecutionStore, "begin", simultaneous_begin)
        with ThreadPoolExecutor(max_workers=2) as workers:
            futures = [
                workers.submit(
                    reprocess_source,
                    root,
                    first.source_id,
                    version=preview["version"],
                    scope=public,
                )
                for _ in range(2)
            ]
            results = [future.result(timeout=30) for future in futures]
        assert sorted(result.status for result in results) == ["added", "blocked"]
        detail = read_knowledge_task(
            root, failed_context.storage_execution.task_id, scope=storage_scope(root, identity)
        )
        assert len(detail.executions) == 2
        assert len(service.requests) == 4
