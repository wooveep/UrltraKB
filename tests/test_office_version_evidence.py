"""Version evidence is anchored to native covers and never inferred from prose mentions."""

from xml.sax.saxutils import escape
from zipfile import ZipFile

import pytest

pytest_plugins = ("test_office_import",)


def docx_cover(path, lines):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with ZipFile(path, "w") as package:
        package.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="{ns}"><w:body>'
            + "".join(f"<w:p><w:r><w:t>{escape(line)}</w:t></w:r></w:p>" for line in lines)
            + "</w:body></w:document>",
        )


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_api_cover_chinese_product_inline_version(tmp_path, fmt):
    from openkb.version_evidence import source_title_candidates

    path = tmp_path / f"unrelated.{fmt}"
    lines = ["开发指南", "云宏WinStack 虚拟化云平台", "日期：2024-03-15   版本：V9.3.1"]
    if fmt == "pdf":
        import pymupdf

        with pymupdf.open() as pdf:
            pdf.new_page().insert_textbox((40, 40, 550, 400), "\n".join(lines), fontname="china-s")
            pdf.save(path)
    else:
        docx_cover(path, lines)
    candidates = source_title_candidates(path, fmt, path.name)
    assert {(c.field, c.values) for c in candidates if c.confidence == "verified"} >= {
        ("product", ("云宏WinStack 虚拟化云平台",)),
        ("family", ("api guide",)),
        ("applicable_versions", ("9.3.1",)),
    }
    assert all(c.location and c.excerpt and c.policy for c in candidates)


def test_xlsx_title_fields_and_filename_hint_do_not_adopt_third_party_version(tmp_path):
    from openpyxl import Workbook

    from openkb.version_evidence import source_title_candidates

    path = tmp_path / "CNware-WinStack-V9.3.1-产品功能列表.xlsx"
    book = Workbook()
    book.active["A1"] = "CNware WinStack 企业版/信创版"
    book.active["C3"] = "支持迁移到 WinSphere V9.3.1"
    book.save(path)
    candidates = source_title_candidates(path, "xlsx", path.name)
    assert any(c.field == "applicable_versions" and c.confidence == "hint" for c in candidates)
    assert not any(
        c.field == "applicable_versions" and c.confidence == "verified" for c in candidates
    )
    book.active["A1"] = "产品：CNware WinStack"
    book.active["A2"] = "适用版本：V9.4.0"
    book.save(path)
    candidates = source_title_candidates(path, "xlsx", path.name)
    assert {
        c.values
        for c in candidates
        if c.field == "applicable_versions" and c.confidence == "verified"
    } == {("9.4.0",)}
    assert any("A2" in c.location for c in candidates if c.confidence == "verified")


def test_pptx_uses_first_slide_in_presentation_order(tmp_path):
    from pptx import Presentation

    from openkb.version_evidence import source_title_candidates

    deck = Presentation()
    for title in ("产品：Wrong\n适用版本：V1.0", "接口开发指南\n产品：Right\n适用版本：V9.4.0"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        from pptx.util import Inches

        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3)).text = title
    ids = deck.slides._sldIdLst
    ids.insert(0, ids[-1])
    path = tmp_path / "unrelated.pptx"
    deck.save(path)
    candidates = source_title_candidates(path, "pptx", path.name)
    assert {
        c.values for c in candidates if c.field == "product" and c.confidence == "verified"
    } == {("Right",)}
    assert {
        c.values
        for c in candidates
        if c.field == "applicable_versions" and c.confidence == "verified"
    } == {("9.4.0",)}


def test_filename_hints_are_not_automatically_applied(kb_dir, tmp_path):
    from openkb.inputs import prepared_input
    from openkb.source_catalog import admit_source_revision
    from openkb.version_metadata import assess_version

    path = tmp_path / "CNware-WinStack-V9.3.1-产品功能列表.docx"
    docx_cover(path, ["支持迁移到 WinSphere V9.3.1"])
    # Admission is enough to test the metadata boundary; no converter/model is needed.
    with prepared_input(path) as prepared:
        admission = admit_source_revision(kb_dir, prepared)
    result = assess_version(kb_dir, admission, None)
    assert not result.metadata.applicable_versions
    assert not result.candidates


def test_explicit_reevaluation_retains_old_knowledge_until_resume(kb_dir, pdf_model, monkeypatch):
    import pymupdf

    from openkb import version_evidence
    from openkb.application.documents import import_document
    from openkb.application.version_evidence import reevaluate_source_version
    from openkb.source_catalog import read_source
    from openkb.unit_publication import read_head

    path = kb_dir / "unrelated.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_textbox(
            (40, 40, 550, 400),
            "开发指南\n云宏WinStack 虚拟化云平台\n日期：2024-03-15   版本：V9.3.1",
            fontname="china-s",
        )
        pdf.save(path)
    current = version_evidence.source_title_candidates
    monkeypatch.setattr(version_evidence, "source_title_candidates", lambda *a, **k: ())
    first = import_document(kb_dir, path)
    source = read_source(kb_dir, first.source_id)
    head = read_head(kb_dir, first.units[0].view_id)
    monkeypatch.setattr(version_evidence, "source_title_candidates", current)
    review = reevaluate_source_version(kb_dir, first.source_id)
    assert review.status == "ready"
    assert review.metadata.applicable_versions == ("9.3.1",)
    assert read_source(kb_dir, first.source_id) == source
    assert read_head(kb_dir, first.units[0].view_id) == head
    assert reevaluate_source_version(kb_dir, first.source_id) == review


def test_reevaluation_drops_obsolete_automatic_evidence_but_keeps_user_fields(
    kb_dir, pdf_model, monkeypatch
):
    import pymupdf

    from openkb.application.documents import import_document
    from openkb.application.version_evidence import reevaluate_source_version
    from openkb.unit_publication import read_head
    from openkb.view_records import SourceMetadata, VersionCandidate

    path = kb_dir / "guide.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Product facts.")
        pdf.save(path)
    candidates = tuple(
        VersionCandidate(field=field, values=value, location="metadata.title", excerpt="Title")
        for field, value in (("product", ("Acme",)), ("applicable_versions", ("1.0",)))
    )
    monkeypatch.setattr("openkb.version_evidence.source_title_candidates", lambda *args: candidates)
    first = import_document(kb_dir, path, metadata=SourceMetadata(family="Confirmed family"))
    assert first.status == "added"
    head = read_head(kb_dir, first.units[0].view_id)
    candidates = tuple(c.model_copy(update={"confidence": "hint"}) for c in candidates)
    review = reevaluate_source_version(kb_dir, first.source_id)
    assert review.status == "blocked"
    assert review.metadata.product is None
    assert review.metadata.applicable_versions == ()
    assert review.metadata.family == "Confirmed family"
    assert read_head(kb_dir, first.units[0].view_id) == head
