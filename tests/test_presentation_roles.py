"""Native presentation roles survive the public import and source-reading boundaries."""

import asyncio
import json

import pytest
from agents.tool_context import ToolContext
from pptx import Presentation
from pptx.util import Inches

from openkb.application.documents import import_document
from openkb.evidence import BlockDraft, ParseStore
from tests.docx_attachment_fixtures import attached_docx
from tests.http_model_fixture import evidence_response


@pytest.mark.parametrize("embedded", [False, True])
@pytest.mark.parametrize("layout, role", [(6, None), (0, "CENTER_TITLE"), (1, "TITLE")])
def test_native_placeholder_role_is_preserved_without_guessing_from_names(
    kb_dir, tmp_path, model_service, layout, role, embedded
):
    from openkb.agent.source_tools import source_tools

    path = tmp_path / "instructions.pptx"
    book = Presentation()
    slide = book.slides.add_slide(book.slide_layouts[layout])
    heading = slide.shapes.title
    if heading is None:
        heading = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(8), Inches(1))
    heading.name = "Title 1"
    heading.text = "Wait for pressure to fall before opening."
    footer = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(8), Inches(0.5))
    footer.text = "Workshop instructions"
    book.save(path)
    title_id = heading.shape_id if role else None
    if embedded:
        path = attached_docx(tmp_path / "parent.docx", path.read_bytes(), name=path.name)

    observed = []

    def respond(body):
        payload = json.loads(body["messages"][-1]["content"])
        if payload["stage"] in {"planning", "generation", "verification"}:
            observed.append(payload)
        return evidence_response(payload)

    model_service.respond = respond
    result = import_document(kb_dir, path)
    assert result.status == "added" and result.knowledge_compilation == "completed", result
    if embedded:
        from openkb.application.attachment_imports import import_attachment
        from openkb.application.execution import ExecutionContext
        from openkb.runtime.requests import ImportAttachment
        from openkb.sources import SourceStore

        assert result.attachments == ()
        assert heading.text not in str(observed)
        observed.clear()
        store = SourceStore(kb_dir)
        parent = store.version(result.input_version)
        assert store.list_sources() == (parent,)
        parsed = ParseStore(kb_dir).load(result.parse_id)
        item = parsed.blocks[0].location["attachment_files"][0]
        child = store.intake_attachment(
            parent,
            part=item["part"],
            name=item["name"],
            content=store.asset(item["blob"]).read_bytes(),
        )
        result = import_attachment(
            kb_dir,
            ImportAttachment(
                child.source_id, child.id, result.input_version, item["part"], child.name
            ),
            context=ExecutionContext(),
        )
        assert result.knowledge_compilation == "completed", result
    parsed = ParseStore(kb_dir).load(result.parse_id)
    block = next(b for b in parsed.blocks if b.location.get("object_id") == heading.shape_id)
    assert block.location["placeholder_type"] == role
    assert block.location["title_object_id"] == title_id
    assert block.location["title_placeholder_count"] == (1 if role else 0)
    assert block.kind == ("heading" if role else "paragraph")
    assert "Slide title: Slide 1" not in block.context
    assert {p["stage"] for p in observed} == {"planning", "generation", "verification"}
    for payload in observed:
        rows = payload["evidence"]["blocks"]
        row = next(r for r in rows if r["text"] == heading.text)
        assert row["location"]["placeholder_type"] == role
        assert row["location"]["title_object_id"] == title_id

    tools = {t.name: t for t in source_tools(kb_dir)[0]}

    def invoke(name, **args):
        return json.loads(
            asyncio.run(
                tools[name].on_invoke_tool(
                    ToolContext(
                        context=None,
                        tool_name=name,
                        tool_call_id="presentation-role",
                        tool_arguments="{}",
                    ),
                    json.dumps(args),
                )
            )
        )

    tree = invoke("read_source_tree", source_id=result.source_id)
    node = next(n for n in tree["nodes"] if n["start"] <= block.order < n["end"])
    response = invoke(
        "read_source_node",
        source_id=result.source_id,
        node_id=node["id"],
        offset=block.order - node["start"],
        max_chars=128,
    )
    row = response["evidence"][0]
    assert "presentation_roles" in response["evidence_provenance"]
    assert row["location"]["placeholder_type"] == role
    assert row["location"]["title_object_id"] == title_id
    assert len(row["text"]) + len(row["context"]) <= 128
    matching = invoke("search_source_text", source_id=result.source_id, query=heading.text)
    assert "presentation_roles" in matching["evidence_provenance"]


@pytest.mark.parametrize(
    "location",
    [
        {"kind": "pptx", "slide": 1, "object_id": 2, "placeholder_type": True},
        {"kind": "pptx", "slide": 1, "object_id": 2, "placeholder_type": "HEADING"},
        {"kind": "pptx", "slide": 1, "placeholder_type": None},
        {
            "kind": "pptx",
            "slide": 1,
            "object_id": 2,
            "title_object_id": False,
            "title_placeholder_count": 1,
        },
        {
            "kind": "pptx",
            "slide": 1,
            "object_id": 2,
            "title_object_id": 0,
            "title_placeholder_count": 1,
        },
        {"kind": "docx", "paragraph": 1, "placeholder_type": None},
        {
            "kind": "pptx",
            "slide": 1,
            "object_id": 2,
            "placeholder_type": "TITLE",
            "title_placeholder_count": 0,
            "title_object_id": None,
        },
        {
            "kind": "pptx",
            "slide": 1,
            "object_id": 2,
            "placeholder_type": None,
            "title_placeholder_count": 1,
            "title_object_id": 2,
        },
    ],
)
def test_native_placeholder_metadata_rejects_invalid_types_and_wrong_source(location):
    with pytest.raises(ValueError):
        BlockDraft("Preserved original wording.", "paragraph", location)


@pytest.mark.parametrize("second_role", ["obj", "title", None])
def test_native_title_uses_its_type_instead_of_placeholder_index_zero(
    kb_dir, tmp_path, model_service, second_role
):
    book = Presentation()
    slide = book.slides.add_slide(book.slide_layouts[1])
    title, body = slide.shapes.title, slide.placeholders[1]
    title.text, body.text = "Stop before opening.", "Plain body content."
    for shape in (title, body):
        shape.left, shape.top, shape.width, shape.height = (
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(1),
        )
    title._element.xpath("./p:nvSpPr/p:nvPr/p:ph")[0].set("idx", "7")
    body._element.xpath("./p:nvSpPr/p:nvPr/p:ph")[0].set("idx", "0")
    if second_role is None:
        body._element.getparent().remove(body._element)
    else:
        body._element.xpath("./p:nvSpPr/p:nvPr/p:ph")[0].set("type", second_role)
    source = tmp_path / "nonstandard-indices.pptx"
    book.save(source)
    imported = import_document(kb_dir, source)
    assert imported.status == "added" and imported.knowledge_compilation == "completed", imported
    parsed = ParseStore(kb_dir).load(imported.parse_id)
    by_object = {block.location["object_id"]: block for block in parsed.blocks}
    assert by_object[title.shape_id].kind == "heading"
    if second_role:
        assert by_object[body.shape_id].kind == (
            "heading" if second_role == "title" else "paragraph"
        )
    for block in parsed.blocks:
        assert block.location["title_placeholder_count"] == (2 if second_role == "title" else 1)
        assert block.location["title_object_id"] == (
            None if second_role == "title" else title.shape_id
        )
        assert "No native title placeholder" not in block.context
