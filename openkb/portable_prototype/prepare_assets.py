"""Build-only fixture/resource assembly. Run using the pinned prototype environment."""

import argparse
import hashlib
import json
import os
import shutil
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--renderer-source", type=Path, required=True)
    args = parser.parse_args()
    source = args.renderer_source.resolve()
    runtime = ROOT / ".runtime"
    renderers = runtime / "renderers"
    renderers.mkdir(parents=True, exist_ok=True)
    suffix = "win-x64/node.exe" if os.name == "nt" else "linux-x64/bin/node"
    extension = ".exe" if os.name == "nt" else ""
    shutil.copy2(source / ".runtime" / f"node-v24.20.0-{suffix}", renderers / f"node{extension}")
    for directory in ("release", "debug"):
        helper = (
            source / "rust-helper" / "target" / directory / f"openkb-render-prototype{extension}"
        )
        if helper.exists():
            shutil.copy2(helper, renderers / f"renderer{extension}")
            break
    else:
        raise FileNotFoundError("Build the renderer helper on this target system first")
    shutil.copy2(source / "mathjax_render.mjs", renderers)
    for start, end in (
        (source / "node_modules", renderers / "node_modules"),
        (source / ".runtime" / "fonts", renderers / "fonts"),
    ):
        shutil.copytree(start, end, dirs_exist_ok=True)

    os.environ["TIKTOKEN_CACHE_DIR"] = str(runtime / "tiktoken-cache")
    import tiktoken

    for encoding in ("cl100k_base", "o200k_base"):
        tiktoken.get_encoding(encoding)

    fixtures = runtime / "fixtures"
    fixtures.mkdir(exist_ok=True)
    from PIL import Image

    Image.new("RGB", (64, 48), (31, 150, 130)).save(fixtures / "示意图.png")
    (fixtures / "中文 笔记.md").write_text(
        "# PortableProbe 中文笔记\n\n正文与相对图片。\n\n![示意](示意图.png)\n", encoding="utf-8"
    )
    (fixtures / "中文 页面.html").write_text(
        "<html><body><article><h1>PortableProbe 中文网页</h1>"
        + "<p>PortableProbe is a local document fixture with enough prose for extraction. " * 12
        + "</p></article></body></html>",
        encoding="utf-8",
    )
    with zipfile.ZipFile(fixtures / "中文 文档.docx", "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.'
            'wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
        )
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>PortableProbe 中文文档</w:t></w:r></w:p></w:body></w:document>",
        )
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "PortableProbe 中文演示"
    slide.placeholders[1].text = "本地 Office 转换样本"
    deck.save(fixtures / "中文 演示.pptx")
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["PortableProbe", "中文表格", 42])
    book.save(fixtures / "中文 表格.xlsx")
    import xlwt

    old = xlwt.Workbook()
    sheet = old.add_sheet("样本")
    sheet.write(0, 0, "PortableProbe 中文旧表格")
    sheet.write(1, 0, 42)
    old.save(str(fixtures / "中文 旧表格.xls"))
    import pymupdf

    for name, count in (("中文 短文.pdf", 2), ("长文.pdf", 20)):
        pdf = pymupdf.open()
        for number in range(count):
            page = pdf.new_page()
            page.insert_font(
                fontname="Noto", fontfile=str(renderers / "fonts" / "NotoSansCJKsc-Regular.otf")
            )
            page.insert_text(
                (40, 70), f"PortableProbe 中文 PDF 第 {number + 1} 页", fontname="Noto"
            )
            page.insert_image(pymupdf.Rect(40, 90, 104, 138), filename=str(fixtures / "示意图.png"))
        pdf.subset_fonts()
        pdf.save(fixtures / name, deflate=True)
        pdf.close()
    manifest = {
        str(p.relative_to(runtime)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(runtime.rglob("*"))
        if p.is_file()
    }
    (runtime / "asset-hashes.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Assembled {len(manifest)} local resource files")


if __name__ == "__main__":
    main()
